from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.mail import send_mail
from django.core.paginator import Paginator
from django.db import IntegrityError, models, transaction
from django.db.models import Count
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.utils.text import slugify
from django.views.decorators.http import require_http_methods
from django.core.exceptions import PermissionDenied

KB_PAGE_SIZE = 10
GLOBAL_SEARCH_LIMIT = 8

from accounts.models import Info

from .forms import (
    ArticleForm,
    ArticleCommentForm,
    QuestionForm,
    AnswerForm,
    ReportForm,
    ReportCommentForm,
    IssueLearningForm,
    IssueLearningCommentForm,
    IssueLearningExcelUploadForm,
)
from .models import (
    Article,
    ArticleAttachment,
    ArticleComment,
    ArticleRevision,
    Category,
    Notification,
    Question,
    Answer,
    Report,
    ReportComment,
    ReportAttachment,
    IssueLearning,
    IssueLearningComment,
    IssueLearningAttachment,
    Tag,
    UserProfile,
    Vote,
)
from .sanitize import sanitize_html


def _notify(recipient, actor, verb, **target):
    """Create an in-app notification. No-op if the actor is the recipient."""
    if recipient is None or actor is None or recipient.pk == actor.pk:
        return
    Notification.objects.create(recipient=recipient, actor=actor, verb=verb, **target)


def _apply_search(queryset, query, high_fields, low_fields=()):
    """
    Text search across every given field, with a two-tier relevance:
    `high_fields` (title, body/description, tags) rank above `low_fields`
    (e.g. category/application name). Adds a `relevance` annotation --
    the caller should order by `-relevance` first, then its usual ordering.

    A field path that crosses an M2M (e.g. 'tags__name') can multiply rows
    when it matches more than one related row per object, so this always
    ends in .distinct() -- callers with a Count() annotation in the same
    queryset should pass distinct=True to it to avoid inflated counts.
    """
    if not query:
        return queryset

    q_filter = models.Q()
    for field in list(high_fields) + list(low_fields):
        q_filter |= models.Q(**{f'{field}__icontains': query})

    relevance = models.Case(
        *[models.When(**{f'{field}__icontains': query}, then=models.Value(2)) for field in high_fields],
        default=models.Value(1),
        output_field=models.IntegerField(),
    )
    return queryset.filter(q_filter).annotate(relevance=relevance).distinct()


def _filter_issues(request):
    """
    Shared Issue & Learning filtering, used by both the paginated list in
    forum_home() and the (unpaginated) Excel export -- so "download the
    issues I'm currently looking at" and "the list I'm currently looking at"
    can never drift apart.
    """
    query = request.GET.get('q', '').strip()
    tag_slug = request.GET.get('tag', '').strip()
    current_tag = Tag.objects.filter(slug=tag_slug).first() if tag_slug else None
    issue_project = request.GET.get('issue_project', 'all')
    issue_status = request.GET.get('issue_status', 'all')
    issue_priority = request.GET.get('issue_priority', 'all')
    issue_capa_status = request.GET.get('issue_capa_status', 'all')

    issues = IssueLearning.objects.select_related('project', 'reporter').prefetch_related(
        'tags', 'comments__author'
    ).annotate(comment_count=Count('comments', distinct=True))
    issues = _apply_search(
        issues, query,
        high_fields=['problem_statement', 'issue_description', 'tags__name', 'closure_accountability'],
        low_fields=['project__project_id', 'location_area'],
    )
    if current_tag:
        issues = issues.filter(tags=current_tag)
    if issue_project != 'all':
        issues = issues.filter(project_id=issue_project)
    if issue_status != 'all':
        issues = issues.filter(status=issue_status)
    if issue_priority != 'all':
        issues = issues.filter(priority=issue_priority)
    if issue_capa_status != 'all':
        issues = issues.filter(capa_status=issue_capa_status)
    issues = issues.order_by('-relevance', '-updated_at') if query else issues.order_by('-updated_at')
    return issues


def _resolve_tags(names):
    """Free-typed tag names -> Tag rows, creating any that don't exist yet."""
    tags = []
    for name in names:
        slug = slugify(name, allow_unicode=True)
        if not slug:
            continue
        try:
            with transaction.atomic():
                tag, _ = Tag.objects.get_or_create(slug=slug, defaults={'name': name})
        except IntegrityError:
            # Someone else created the same brand-new tag in the instant
            # between our lookup and our insert -- just use theirs.
            tag = Tag.objects.get(slug=slug)
        tags.append(tag)
    return tags


def _finalize_richtext(instance, form, field_name, format_field_name=None, is_html_attr=None):
    """
    Apply the optional rich-text toggle's result to a just-populated (but
    not yet saved) model instance: sanitize and flag as HTML if the user
    turned formatting on, otherwise leave the plain text as-is.
    """
    format_field_name = format_field_name or f'{field_name}_format'
    is_html_attr = is_html_attr or ('is_html' if field_name == 'body' else f'{field_name}_is_html')
    is_html = form.cleaned_data.get(format_field_name) == 'html'
    value = getattr(instance, field_name)
    if is_html:
        value = sanitize_html(value)
    setattr(instance, field_name, value)
    setattr(instance, is_html_attr, is_html)


@login_required
def home(request):
    return render(request, 'base.html')


@login_required
def downloads(request):
    infos = Info.objects.all()
    context = {
        'infos': infos
    }
    return render(request, 'home/downloads.html', context)


def superuser_required(view_func):
    def _wrapped_view(request, *args, **kwargs):
        if not request.user.is_authenticated or not request.user.is_superuser:
            raise PermissionDenied
        return view_func(request, *args, **kwargs)
    return _wrapped_view


@superuser_required
def clear_cache(request):
    cache.clear()
    return JsonResponse({"status": "ok", "message": "Cache cleared"})


def build_article_hierarchy(articles):
    articles_list = list(articles)
    if not articles_list:
        return []

    articles_by_parent = {}
    articles_by_id = {}
    for article in articles_list:
        articles_by_id[article.id] = article
        articles_by_parent.setdefault(article.parent_id, []).append(article)

    roots = [article for article in articles_list if article.is_hierarchy_root]
    if not roots:
        roots = [article for article in articles_list if article.parent_id is None]

    for article in articles_list:
        if article.is_hierarchy_root:
            continue
        if article.parent_id and article.parent_id not in articles_by_id:
            roots.append(article)

    unique_roots = []
    seen_root_ids = set()
    for root in roots:
        if root.id in seen_root_ids:
            continue
        seen_root_ids.add(root.id)
        unique_roots.append(root)
    roots = sorted(unique_roots, key=lambda item: item.title.lower())

    hierarchy = []
    visited = set()

    def add_node(node, depth):
        if node.id in visited:
            return
        visited.add(node.id)
        hierarchy.append({'article': node, 'indent': depth * 18})
        children = articles_by_parent.get(node.id, [])
        for child in sorted(children, key=lambda item: item.title.lower()):
            add_node(child, depth + 1)

    for root in roots:
        add_node(root, 0)

    added_ids = {item['article'].id for item in hierarchy}
    leftovers = [article for article in articles_list if article.id not in added_ids]
    for article in sorted(leftovers, key=lambda item: item.title.lower()):
        hierarchy.append({'article': article, 'indent': 0})

    return hierarchy


def build_article_tree(articles):
    articles_list = list(articles)
    if not articles_list:
        return []

    articles_by_parent = {}
    articles_by_id = {}
    for article in articles_list:
        articles_by_id[article.id] = article
        articles_by_parent.setdefault(article.parent_id, []).append(article)

    roots = [article for article in articles_list if article.is_hierarchy_root]
    if not roots:
        roots = [article for article in articles_list if article.parent_id is None]

    for article in articles_list:
        if article.is_hierarchy_root:
            continue
        if article.parent_id and article.parent_id not in articles_by_id:
            roots.append(article)

    unique_roots = []
    seen_root_ids = set()
    for root in roots:
        if root.id in seen_root_ids:
            continue
        seen_root_ids.add(root.id)
        unique_roots.append(root)
    roots = sorted(unique_roots, key=lambda item: item.title.lower())

    visited = set()

    def build_node(node):
        if node.id in visited:
            return None
        visited.add(node.id)
        children = []
        for child in sorted(articles_by_parent.get(node.id, []), key=lambda item: item.title.lower()):
            child_node = build_node(child)
            if child_node:
                children.append(child_node)
        return {'article': node, 'children': children}

    tree = []
    for root in roots:
        node = build_node(root)
        if node:
            tree.append(node)

    leftovers = [article for article in articles_list if article.id not in visited]
    for article in sorted(leftovers, key=lambda item: item.title.lower()):
        tree.append({'article': article, 'children': []})

    return tree


def get_kb_stats_context():
    total_articles = Article.objects.count()
    total_questions = Question.objects.count()
    total_reports = Report.objects.count()
    total_interactions = Answer.objects.count() + ReportComment.objects.count()
    total_users = get_user_model().objects.count()
    open_bugs = Report.objects.filter(
        type=Report.TYPE_BUG,
        status__in=[Report.STATUS_OPEN, Report.STATUS_IN_PROGRESS]
    ).count()
    open_features = Report.objects.filter(
        type=Report.TYPE_FEATURE,
        status__in=[Report.STATUS_OPEN, Report.STATUS_IN_PROGRESS]
    ).count()
    total_issues = IssueLearning.objects.count()
    open_issues = IssueLearning.objects.filter(
        status__in=[IssueLearning.STATUS_OPEN, IssueLearning.STATUS_UNDER_OBSERVATION]
    ).count()
    closed_issues = IssueLearning.objects.filter(status=IssueLearning.STATUS_CLOSED).count()
    capa_open = IssueLearning.objects.filter(
        capa_status__in=[IssueLearning.CAPA_STATUS_OPEN, IssueLearning.CAPA_STATUS_IN_PROGRESS]
    ).count()
    capa_closed = IssueLearning.objects.filter(
        capa_status__in=[IssueLearning.CAPA_STATUS_COMPLETED, IssueLearning.CAPA_STATUS_VERIFIED]
    ).count()
    return {
        'total_articles': total_articles,
        'total_questions': total_questions,
        'total_reports': total_reports,
        'total_interactions': total_interactions,
        'total_users': total_users,
        'open_bugs': open_bugs,
        'open_features': open_features,
        'total_issues': total_issues,
        'open_issues': open_issues,
        'closed_issues': closed_issues,
        'capa_open': capa_open,
        'capa_closed': capa_closed,
    }


def get_kb_sidebar_context():
    hierarchy_tree = build_article_tree(
        Article.objects.only('id', 'title', 'slug', 'parent_id', 'is_hierarchy_root')
        .filter(models.Q(is_hierarchy_root=True) | models.Q(parent__isnull=False))
        .order_by('title')
    )
    other_articles_sidebar = Article.objects.only('id', 'title', 'slug').filter(
        parent__isnull=True,
        is_hierarchy_root=False
    ).order_by('title')
    return {
        'hierarchy_tree': hierarchy_tree,
        'other_articles_sidebar': other_articles_sidebar,
    }


@login_required
def forum_home(request):
    active_tab = request.GET.get('tab', 'wiki')
    if active_tab not in {'wiki', 'qa', 'reports', 'issues'}:
        active_tab = 'wiki'
    query = request.GET.get('q', '').strip()
    tag_slug = request.GET.get('tag', '').strip()
    current_tag = Tag.objects.filter(slug=tag_slug).first() if tag_slug else None

    context = {
        'active_tab': active_tab,
        'query': query,
        'current_tag': current_tag,
    }
    context.update(get_kb_stats_context())

    # Global search: the header search box submits `q` with no `tab`, so a
    # search from anywhere in the KB looks across every content type at
    # once, instead of only whichever tab happened to be open. Tab links,
    # sidebar filters, and "view all" links all pass `tab` explicitly, which
    # opts back into the single-type, paginated, filterable view below.
    if query and 'tab' not in request.GET:
        wiki_matches = _apply_search(
            Article.objects.select_related('category', 'author').prefetch_related('tags'),
            query, high_fields=['title', 'content', 'excerpt', 'tags__name'], low_fields=['category__name'],
        ).order_by('-relevance', '-updated_at')
        qa_matches = _apply_search(
            Question.objects.select_related('author').prefetch_related('tags', 'tagged_users').annotate(
                answer_count=Count('answers', distinct=True)
            ),
            query, high_fields=['title', 'body', 'tags__name'],
        ).order_by('-relevance', '-created_at')
        report_matches = _apply_search(
            Report.objects.select_related('application', 'reporter').prefetch_related('tags').annotate(
                comment_count=Count('comments', distinct=True)
            ),
            query, high_fields=['title', 'description', 'tags__name'], low_fields=['application__name'],
        ).order_by('-relevance', '-updated_at')
        issue_matches = _apply_search(
            IssueLearning.objects.select_related('project', 'reporter').prefetch_related('tags').annotate(
                comment_count=Count('comments', distinct=True)
            ),
            query,
            high_fields=['problem_statement', 'issue_description', 'tags__name', 'closure_accountability'],
            low_fields=['project__project_id', 'location_area'],
        ).order_by('-relevance', '-updated_at')

        context.update(get_kb_sidebar_context())
        context.update({
            'is_global_search': True,
            'search_wiki_results': wiki_matches[:GLOBAL_SEARCH_LIMIT],
            'search_wiki_total': wiki_matches.count(),
            'search_qa_results': qa_matches[:GLOBAL_SEARCH_LIMIT],
            'search_qa_total': qa_matches.count(),
            'search_report_results': report_matches[:GLOBAL_SEARCH_LIMIT],
            'search_report_total': report_matches.count(),
            'search_issue_results': issue_matches[:GLOBAL_SEARCH_LIMIT],
            'search_issue_total': issue_matches.count(),
        })
        return render(request, 'home/forum_home.html', context)

    if active_tab == 'wiki':
        articles = Article.objects.select_related('category', 'author', 'parent').prefetch_related('tags')
        articles = _apply_search(
            articles, query, high_fields=['title', 'content', 'excerpt', 'tags__name'], low_fields=['category__name'],
        )
        if current_tag:
            articles = articles.filter(tags=current_tag)
        articles = articles.order_by('-relevance', '-updated_at') if query else articles.order_by('-updated_at')

        if query or current_tag:
            # Flat, paginated results read better than a hierarchy built from a partial match set.
            hierarchy_articles = []
            other_articles_qs = articles
        else:
            hierarchy_articles = build_article_hierarchy(
                articles.filter(models.Q(is_hierarchy_root=True) | models.Q(parent__isnull=False))
            )
            other_articles_qs = articles.filter(parent__isnull=True, is_hierarchy_root=False)

        page_obj = Paginator(other_articles_qs, KB_PAGE_SIZE).get_page(request.GET.get('page'))
        context.update(get_kb_sidebar_context())
        context.update({
            'hierarchy_articles': hierarchy_articles,
            'other_articles': page_obj,
            'page_obj': page_obj,
        })
    elif active_tab == 'qa':
        qa_status = request.GET.get('qa_status', 'all')
        questions = Question.objects.select_related('author', 'accepted_answer').prefetch_related('tags', 'tagged_users').annotate(
            answer_count=Count('answers', distinct=True)
        )
        questions = _apply_search(questions, query, high_fields=['title', 'body', 'tags__name'])
        if current_tag:
            questions = questions.filter(tags=current_tag)
        if qa_status == 'solved':
            questions = questions.filter(is_solved=True)
        elif qa_status == 'unsolved':
            questions = questions.filter(is_solved=False)
        elif qa_status == 'tagged_to_me':
            questions = questions.filter(tagged_users=request.user)
        questions = questions.order_by('-relevance', '-created_at') if query else questions.order_by('-created_at')

        page_obj = Paginator(questions, KB_PAGE_SIZE).get_page(request.GET.get('page'))
        context.update({
            'qa_status': qa_status,
            'questions': page_obj,
            'page_obj': page_obj,
        })
    elif active_tab == 'reports':
        report_type = request.GET.get('report_type', 'all')
        report_status = request.GET.get('report_status', 'all')
        report_priority = request.GET.get('report_priority', 'all')
        reports = Report.objects.select_related('application', 'reporter', 'assignee').prefetch_related('tags').annotate(
            comment_count=Count('comments', distinct=True)
        )
        reports = _apply_search(
            reports, query, high_fields=['title', 'description', 'tags__name'], low_fields=['application__name'],
        )
        if current_tag:
            reports = reports.filter(tags=current_tag)
        if report_type != 'all':
            reports = reports.filter(type=report_type)
        if report_status != 'all':
            reports = reports.filter(status=report_status)
        if report_priority != 'all':
            reports = reports.filter(priority=report_priority)
        reports = reports.order_by('-relevance', '-updated_at') if query else reports.order_by('-updated_at')

        page_obj = Paginator(reports, KB_PAGE_SIZE).get_page(request.GET.get('page'))
        context.update({
            'report_type': report_type,
            'report_status': report_status,
            'report_priority': report_priority,
            'reports': page_obj,
            'page_obj': page_obj,
        })
    else:
        issue_project = request.GET.get('issue_project', 'all')
        issue_status = request.GET.get('issue_status', 'all')
        issue_priority = request.GET.get('issue_priority', 'all')
        issue_capa_status = request.GET.get('issue_capa_status', 'all')
        issues = _filter_issues(request)

        from planner.models import Project as PlannerProject
        page_obj = Paginator(issues, KB_PAGE_SIZE).get_page(request.GET.get('page'))
        context.update({
            'issue_project': issue_project,
            'issue_status': issue_status,
            'issue_priority': issue_priority,
            'issue_capa_status': issue_capa_status,
            'issue_projects': PlannerProject.objects.filter(kb_issues__isnull=False).distinct().order_by('project_id'),
            'issues': page_obj,
            'page_obj': page_obj,
        })

    return render(request, 'home/forum_home.html', context)


@login_required
def tag_detail(request, slug):
    tag = get_object_or_404(Tag, slug=slug)
    wiki_matches = Article.objects.filter(tags=tag).select_related('category', 'author').prefetch_related('tags').order_by('-updated_at')
    qa_matches = Question.objects.filter(tags=tag).select_related('author').prefetch_related('tags', 'tagged_users').annotate(
        answer_count=Count('answers')
    ).order_by('-created_at')
    report_matches = Report.objects.filter(tags=tag).select_related('application', 'reporter').prefetch_related('tags').annotate(
        comment_count=Count('comments')
    ).order_by('-updated_at')
    issue_matches = IssueLearning.objects.filter(tags=tag).select_related('project', 'reporter').prefetch_related('tags').annotate(
        comment_count=Count('comments')
    ).order_by('-updated_at')

    context = {
        'tag': tag,
        'active_tab': 'wiki',
        'search_wiki_results': wiki_matches[:GLOBAL_SEARCH_LIMIT],
        'search_wiki_total': wiki_matches.count(),
        'search_qa_results': qa_matches[:GLOBAL_SEARCH_LIMIT],
        'search_qa_total': qa_matches.count(),
        'search_report_results': report_matches[:GLOBAL_SEARCH_LIMIT],
        'search_report_total': report_matches.count(),
        'search_issue_results': issue_matches[:GLOBAL_SEARCH_LIMIT],
        'search_issue_total': issue_matches.count(),
    }
    context.update(get_kb_stats_context())
    context.update(get_kb_sidebar_context())
    return render(request, 'home/kb_tag_detail.html', context)


@login_required
def article_category(request, slug):
    category = get_object_or_404(Category, slug=slug)
    query = request.GET.get('q', '').strip()
    articles = Article.objects.filter(category=category).select_related('author', 'parent').prefetch_related('tags')
    articles = _apply_search(articles, query, high_fields=['title', 'content', 'excerpt', 'tags__name'])
    articles = articles.order_by('-relevance', '-updated_at') if query else articles.order_by('-updated_at')

    if query:
        hierarchy_articles = []
        other_articles_qs = articles
    else:
        hierarchy_articles = build_article_hierarchy(
            articles.filter(models.Q(is_hierarchy_root=True) | models.Q(parent__isnull=False))
        )
        other_articles_qs = articles.filter(parent__isnull=True, is_hierarchy_root=False)

    page_obj = Paginator(other_articles_qs, KB_PAGE_SIZE).get_page(request.GET.get('page'))
    context = {
        'category': category,
        'query': query,
        'hierarchy_articles': hierarchy_articles,
        'other_articles': page_obj,
        'page_obj': page_obj,
    }
    return render(request, 'home/kb_category.html', context)


@login_required
def user_profile(request, username):
    profile_user = get_object_or_404(get_user_model(), username=username, is_active=True)
    articles = Article.objects.filter(author=profile_user).select_related('category').order_by('-updated_at')
    questions = Question.objects.filter(author=profile_user).annotate(
        answer_count=Count('answers', distinct=True)
    ).order_by('-created_at')
    answers = Answer.objects.filter(author=profile_user).select_related('question').order_by('-created_at')
    reports = Report.objects.filter(reporter=profile_user).select_related('application').order_by('-updated_at')

    try:
        home_profile = profile_user.home_profile
    except UserProfile.DoesNotExist:
        home_profile = None

    context = {
        'profile_user': profile_user,
        'home_profile': home_profile,
        'articles': articles[:GLOBAL_SEARCH_LIMIT],
        'articles_total': articles.count(),
        'questions': questions[:GLOBAL_SEARCH_LIMIT],
        'questions_total': questions.count(),
        'answers': answers[:GLOBAL_SEARCH_LIMIT],
        'answers_total': answers.count(),
        'reports': reports[:GLOBAL_SEARCH_LIMIT],
        'reports_total': reports.count(),
    }
    return render(request, 'home/kb_profile.html', context)


@login_required
def article_detail(request, slug):
    article = get_object_or_404(
        Article.objects.select_related('category', 'author').prefetch_related('tags', 'comments__author', 'attachments'),
        slug=slug
    )

    if request.method == 'POST':
        form = ArticleCommentForm(request.POST)
        if form.is_valid():
            comment = form.save(commit=False)
            comment.article = article
            comment.author = request.user
            _finalize_richtext(comment, form, 'body')
            comment.save()
            _notify(article.author, request.user, 'commented on', article=article)
            messages.success(request, 'Comment added.')
            return redirect(article.get_absolute_url())
        messages.error(request, 'Please correct the errors below.')
    else:
        form = ArticleCommentForm()

    Article.objects.filter(pk=article.pk).update(views=models.F('views') + 1)
    article.refresh_from_db(fields=['views'])
    context = {
        'article': article,
        'comments': article.comments.all(),
        'attachments': article.attachments.all(),
        'revisions': article.revisions.select_related('edited_by')[:20],
        'comment_form': form,
        'active_tab': 'wiki',
    }
    context.update(get_kb_stats_context())
    context.update(get_kb_sidebar_context())
    return render(request, 'home/kb_article_detail.html', context)


@login_required
def article_comment_update(request, pk):
    comment = get_object_or_404(ArticleComment.objects.select_related('article'), pk=pk)
    if not (request.user.is_staff or comment.author == request.user):
        messages.error(request, 'You do not have permission to edit this comment.')
        return redirect(comment.article.get_absolute_url())

    form = ArticleCommentForm(request.POST or None, instance=comment)
    if request.method == 'POST' and form.is_valid():
        comment = form.save(commit=False)
        _finalize_richtext(comment, form, 'body')
        comment.save()
        messages.success(request, 'Comment updated.')
        return redirect(comment.article.get_absolute_url())

    context = {'form': form, 'heading': 'Edit comment', 'back_url': comment.article.get_absolute_url()}
    return render(request, 'home/kb_content_edit.html', context)


@login_required
@require_http_methods(['POST'])
def article_comment_delete(request, pk):
    comment = get_object_or_404(ArticleComment.objects.select_related('article'), pk=pk)
    if not (request.user.is_staff or comment.author == request.user):
        messages.error(request, 'You do not have permission to delete this comment.')
        return redirect(comment.article.get_absolute_url())
    article = comment.article
    comment.delete()
    messages.success(request, 'Comment deleted.')
    return redirect(article.get_absolute_url())


def _notify_tagged_users(request, question, users):
    """In-app notification + best-effort email for people newly tagged on a support request."""
    users = list(users)
    for user in users:
        _notify(user, question.author, 'tagged you on', question=question)

    recipients = [u.email for u in users if u.email and u.id != question.author_id]
    if not recipients:
        return
    url = request.build_absolute_uri(question.get_absolute_url())
    requester_name = question.author.get_full_name() or question.author.username
    html_message = render_to_string('home/email/question_tagged.html', {
        'question': question,
        'requester_name': requester_name,
        'url': url,
    })
    send_mail(
        subject=f"You were tagged on a support request: {question.title}",
        message=f"{requester_name} tagged you on a support request: {question.title}\n\n{url}",
        from_email=None,
        recipient_list=recipients,
        html_message=html_message,
        fail_silently=True,
    )


def _cast_vote(user, target, direction):
    """Toggle/switch a user's vote on a Question or Answer, keeping its `votes` counter in sync."""
    value = Vote.UP if direction == 'up' else Vote.DOWN
    lookup = {'user': user}
    if isinstance(target, Question):
        lookup['question'] = target
    else:
        lookup['answer'] = target

    existing = Vote.objects.filter(**lookup).first()
    if existing is None:
        Vote.objects.create(value=value, **lookup)
        delta = value
    elif existing.value == value:
        existing.delete()
        delta = -value
    else:
        delta = value - existing.value
        existing.value = value
        existing.save(update_fields=['value'])

    if delta:
        type(target).objects.filter(pk=target.pk).update(votes=models.F('votes') + delta)
        target.refresh_from_db(fields=['votes'])
    return target.votes


@login_required
def question_detail(request, pk):
    question = get_object_or_404(
        Question.objects.select_related('author', 'accepted_answer').prefetch_related('tags', 'tagged_users', 'answers__author'),
        pk=pk
    )

    if request.method == 'POST':
        form = AnswerForm(request.POST)
        if form.is_valid():
            answer = form.save(commit=False)
            answer.question = question
            answer.author = request.user
            _finalize_richtext(answer, form, 'body')
            answer.save()
            _notify(question.author, request.user, 'contributed to', question=question)
            messages.success(request, 'Answer posted.')
            return redirect('kb-question-detail', pk=question.pk)
        messages.error(request, 'Please correct the errors below.')
    else:
        form = AnswerForm()

    user_votes = {
        vote.answer_id: vote.value
        for vote in Vote.objects.filter(user=request.user, answer__question=question)
    }
    question_vote = Vote.objects.filter(user=request.user, question=question).first()

    answers = list(question.answers.all())
    for answer in answers:
        answer.user_vote = user_votes.get(answer.id, 0)

    context = {
        'question': question,
        'answers': answers,
        'answer_form': form,
        'active_tab': 'qa',
        'user_question_vote': question_vote.value if question_vote else 0,
        'can_accept_answer': request.user.is_staff or question.author_id == request.user.id,
    }
    context.update(get_kb_stats_context())
    context.update(get_kb_sidebar_context())
    return render(request, 'home/kb_question_detail.html', context)


@login_required
@require_http_methods(['POST'])
def question_vote(request, pk, direction):
    if direction not in ('up', 'down'):
        raise Http404
    question = get_object_or_404(Question, pk=pk)
    if question.author_id == request.user.id:
        messages.error(request, "You can't vote on your own question.")
    else:
        _cast_vote(request.user, question, direction)
    return redirect('kb-question-detail', pk=question.pk)


@login_required
@require_http_methods(['POST'])
def answer_vote(request, pk, direction):
    if direction not in ('up', 'down'):
        raise Http404
    answer = get_object_or_404(Answer.objects.select_related('question'), pk=pk)
    if answer.author_id == request.user.id:
        messages.error(request, "You can't vote on your own answer.")
    else:
        _cast_vote(request.user, answer, direction)
    return redirect('kb-question-detail', pk=answer.question_id)


@login_required
@require_http_methods(['POST'])
def answer_accept(request, pk):
    answer = get_object_or_404(Answer.objects.select_related('question'), pk=pk)
    question = answer.question
    if not (request.user.is_staff or question.author_id == request.user.id):
        messages.error(request, 'Only the question author can accept an answer.')
        return redirect('kb-question-detail', pk=question.pk)

    if question.accepted_answer_id == answer.pk:
        Answer.objects.filter(pk=answer.pk).update(is_accepted=False)
        question.accepted_answer = None
        question.is_solved = False
        question.save(update_fields=['accepted_answer', 'is_solved'])
        messages.success(request, 'Answer unmarked as accepted.')
    else:
        Answer.objects.filter(question=question).update(is_accepted=False)
        Answer.objects.filter(pk=answer.pk).update(is_accepted=True)
        question.accepted_answer = answer
        question.is_solved = True
        question.save(update_fields=['accepted_answer', 'is_solved'])
        _notify(answer.author, request.user, 'accepted your answer to', question=question)
        messages.success(request, 'Answer marked as accepted.')
    return redirect('kb-question-detail', pk=question.pk)


@login_required
def answer_update(request, pk):
    answer = get_object_or_404(Answer.objects.select_related('question'), pk=pk)
    if not (request.user.is_staff or answer.author == request.user):
        messages.error(request, 'You do not have permission to edit this contribution.')
        return redirect(answer.question.get_absolute_url())

    form = AnswerForm(request.POST or None, instance=answer)
    if request.method == 'POST' and form.is_valid():
        answer = form.save(commit=False)
        _finalize_richtext(answer, form, 'body')
        answer.save()
        messages.success(request, 'Contribution updated.')
        return redirect(answer.question.get_absolute_url())

    context = {'form': form, 'heading': 'Edit contribution', 'back_url': answer.question.get_absolute_url()}
    return render(request, 'home/kb_content_edit.html', context)


@login_required
@require_http_methods(['POST'])
def answer_delete(request, pk):
    answer = get_object_or_404(Answer.objects.select_related('question'), pk=pk)
    if not (request.user.is_staff or answer.author == request.user):
        messages.error(request, 'You do not have permission to delete this contribution.')
        return redirect(answer.question.get_absolute_url())

    question = answer.question
    was_accepted = question.accepted_answer_id == answer.pk
    answer.delete()
    if was_accepted:
        question.accepted_answer = None
        question.is_solved = False
        question.save(update_fields=['accepted_answer', 'is_solved'])
    messages.success(request, 'Contribution deleted.')
    return redirect(question.get_absolute_url())


@login_required
def report_detail(request, pk):
    report = get_object_or_404(
        Report.objects.select_related('application', 'reporter', 'assignee').prefetch_related('tags', 'comments__author', 'attachments'),
        pk=pk
    )

    if request.method == 'POST':
        form = ReportCommentForm(request.POST)
        if form.is_valid():
            comment = form.save(commit=False)
            comment.report = report
            comment.author = request.user
            _finalize_richtext(comment, form, 'body')
            comment.save()
            _notify(report.reporter, request.user, 'commented on', report=report)
            messages.success(request, 'Comment added.')
            return redirect('kb-report-detail', pk=report.pk)
        messages.error(request, 'Please correct the errors below.')
    else:
        form = ReportCommentForm()

    context = {
        'report': report,
        'comments': report.comments.all(),
        'attachments': report.attachments.all(),
        'comment_form': form,
        'active_tab': 'reports',
    }
    context.update(get_kb_stats_context())
    context.update(get_kb_sidebar_context())
    return render(request, 'home/kb_report_detail.html', context)


@login_required
def report_comment_update(request, pk):
    comment = get_object_or_404(ReportComment.objects.select_related('report'), pk=pk)
    if not (request.user.is_staff or comment.author == request.user):
        messages.error(request, 'You do not have permission to edit this comment.')
        return redirect(comment.report.get_absolute_url())

    form = ReportCommentForm(request.POST or None, instance=comment)
    if request.method == 'POST' and form.is_valid():
        comment = form.save(commit=False)
        _finalize_richtext(comment, form, 'body')
        comment.save()
        messages.success(request, 'Comment updated.')
        return redirect(comment.report.get_absolute_url())

    context = {'form': form, 'heading': 'Edit comment', 'back_url': comment.report.get_absolute_url()}
    return render(request, 'home/kb_content_edit.html', context)


@login_required
@require_http_methods(['POST'])
def report_comment_delete(request, pk):
    comment = get_object_or_404(ReportComment.objects.select_related('report'), pk=pk)
    if not (request.user.is_staff or comment.author == request.user):
        messages.error(request, 'You do not have permission to delete this comment.')
        return redirect(comment.report.get_absolute_url())
    report = comment.report
    comment.delete()
    messages.success(request, 'Comment deleted.')
    return redirect(report.get_absolute_url())


@login_required
def issue_detail(request, pk):
    issue = get_object_or_404(
        IssueLearning.objects.select_related('project', 'reporter').prefetch_related('tags', 'comments__author', 'attachments'),
        pk=pk
    )

    if request.method == 'POST':
        form = IssueLearningCommentForm(request.POST)
        if form.is_valid():
            comment = form.save(commit=False)
            comment.issue = issue
            comment.author = request.user
            _finalize_richtext(comment, form, 'body')
            comment.save()
            _notify(issue.reporter, request.user, 'commented on', issue=issue)
            messages.success(request, 'Note added.')
            return redirect('kb-issue-detail', pk=issue.pk)
        messages.error(request, 'Please correct the errors below.')
    else:
        form = IssueLearningCommentForm()

    context = {
        'issue': issue,
        'comments': issue.comments.all(),
        'attachments': issue.attachments.all(),
        'comment_form': form,
        'active_tab': 'issues',
    }
    context.update(get_kb_stats_context())
    context.update(get_kb_sidebar_context())
    return render(request, 'home/kb_issue_detail.html', context)


@login_required
def issue_comment_update(request, pk):
    comment = get_object_or_404(IssueLearningComment.objects.select_related('issue'), pk=pk)
    if not (request.user.is_staff or comment.author == request.user):
        messages.error(request, 'You do not have permission to edit this note.')
        return redirect(comment.issue.get_absolute_url())

    form = IssueLearningCommentForm(request.POST or None, instance=comment)
    if request.method == 'POST' and form.is_valid():
        comment = form.save(commit=False)
        _finalize_richtext(comment, form, 'body')
        comment.save()
        messages.success(request, 'Note updated.')
        return redirect(comment.issue.get_absolute_url())

    context = {'form': form, 'heading': 'Edit note', 'back_url': comment.issue.get_absolute_url()}
    return render(request, 'home/kb_content_edit.html', context)


@login_required
@require_http_methods(['POST'])
def issue_comment_delete(request, pk):
    comment = get_object_or_404(IssueLearningComment.objects.select_related('issue'), pk=pk)
    if not (request.user.is_staff or comment.author == request.user):
        messages.error(request, 'You do not have permission to delete this note.')
        return redirect(comment.issue.get_absolute_url())
    issue = comment.issue
    comment.delete()
    messages.success(request, 'Note deleted.')
    return redirect(issue.get_absolute_url())


@login_required
def kb_create(request):
    content_type = request.GET.get('type', '')
    form = None
    template = 'home/kb_create.html'

    if request.method == 'POST':
        content_type = request.POST.get('content_type', content_type)

    if content_type not in {'wiki', 'qa', 'report', 'issue'}:
        content_type = ''

    if content_type == 'wiki':
        form = ArticleForm(request.POST or None, request.FILES or None, user=request.user)
        if request.method == 'POST' and form.is_valid():
            article = form.save(commit=False)
            article.author = request.user
            article.save()
            form.save_m2m()
            article.tags.set(_resolve_tags(form.cleaned_data['tags']))
            attachments = form.cleaned_data.get('attachments', [])
            for uploaded_file in attachments:
                ArticleAttachment.objects.create(
                    article=article,
                    file=uploaded_file,
                    uploaded_by=request.user
                )
            messages.success(request, 'Article created.')
            return redirect(article.get_absolute_url())
        if request.method == 'POST' and form.errors:
            messages.error(request, 'Please correct the errors below.')
    elif content_type == 'qa':
        form = QuestionForm(request.POST or None)
        if request.method == 'POST' and form.is_valid():
            question = form.save(commit=False)
            question.author = request.user
            _finalize_richtext(question, form, 'body')
            already_solved = form.cleaned_data.get('already_solved')
            if already_solved:
                question.is_solved = True
            question.save()
            form.save_m2m()
            question.tags.set(_resolve_tags(form.cleaned_data['tags']))
            question.tagged_users.set(form.cleaned_data['tagged_users'])

            solution = form.cleaned_data.get('solution', '').strip()
            if already_solved and solution:
                solution_is_html = form.cleaned_data.get('solution_format') == 'html'
                answer = Answer.objects.create(
                    question=question,
                    body=sanitize_html(solution) if solution_is_html else solution,
                    is_html=solution_is_html,
                    author=request.user,
                    is_accepted=True,
                )
                question.accepted_answer = answer
                question.save(update_fields=['accepted_answer'])

            _notify_tagged_users(request, question, question.tagged_users.all())
            messages.success(request, 'Support request posted.')
            return redirect(question.get_absolute_url())
        if request.method == 'POST' and form.errors:
            messages.error(request, 'Please correct the errors below.')
    elif content_type == 'report':
        form = ReportForm(request.POST or None, request.FILES or None)
        if request.method == 'POST' and form.is_valid():
            report = form.save(commit=False)
            report.reporter = request.user
            _finalize_richtext(report, form, 'description')
            report.save()
            form.save_m2m()
            report.tags.set(_resolve_tags(form.cleaned_data['tags']))
            attachments = form.cleaned_data.get('attachments', [])
            for uploaded_file in attachments:
                ReportAttachment.objects.create(
                    report=report,
                    file=uploaded_file,
                    uploaded_by=request.user
                )
            messages.success(request, 'Report submitted.')
            return redirect(report.get_absolute_url())
        if request.method == 'POST' and form.errors:
            messages.error(request, 'Please correct the errors below.')
    elif content_type == 'issue':
        form = IssueLearningForm(request.POST or None, request.FILES or None)
        if request.method == 'POST' and form.is_valid():
            issue = form.save(commit=False)
            issue.reporter = request.user
            _finalize_richtext(issue, form, 'issue_description')
            _finalize_richtext(issue, form, 'final_corrective_action')
            issue.save()
            form.save_m2m()
            issue.tags.set(_resolve_tags(form.cleaned_data['tags']))
            attachments = form.cleaned_data.get('attachments', [])
            for uploaded_file in attachments:
                IssueLearningAttachment.objects.create(
                    issue=issue,
                    file=uploaded_file,
                    uploaded_by=request.user
                )
            messages.success(request, 'Issue & learning logged.')
            return redirect(issue.get_absolute_url())
        if request.method == 'POST' and form.errors:
            messages.error(request, 'Please correct the errors below.')

    context = {
        'content_type': content_type,
        'form': form,
    }
    return render(request, template, context)


@login_required
def article_update(request, slug):
    article = get_object_or_404(Article, slug=slug)
    if not (request.user.is_staff or article.author == request.user):
        messages.error(request, 'You do not have permission to edit this article.')
        return redirect(article.get_absolute_url())

    # Captured before the form binds, since Django's ModelForm mutates the
    # instance in place during is_valid() -- reading article.* afterwards
    # would already reflect the *new* values, not what's being replaced.
    original = {'title': article.title, 'content': article.content, 'excerpt': article.excerpt}

    form = ArticleForm(request.POST or None, instance=article, user=request.user)
    if request.method == 'POST' and form.is_valid():
        if set(form.changed_data) & {'title', 'content', 'excerpt'}:
            ArticleRevision.objects.create(article=article, edited_by=request.user, **original)
        form.save()
        article.tags.set(_resolve_tags(form.cleaned_data['tags']))
        messages.success(request, 'Article updated.')
        return redirect(article.get_absolute_url())

    context = {
        'content_type': 'wiki',
        'form': form,
        'object': article,
    }
    return render(request, 'home/kb_create.html', context)


@login_required
@require_http_methods(['POST'])
def article_delete(request, slug):
    article = get_object_or_404(Article, slug=slug)
    if not (request.user.is_staff or article.author == request.user):
        messages.error(request, 'You do not have permission to delete this article.')
        return redirect(article.get_absolute_url())
    article.delete()
    messages.success(request, 'Article deleted.')
    return redirect(f"{reverse('forum-home')}?tab=wiki")


@login_required
@require_http_methods(['POST'])
def article_revision_restore(request, slug, revision_pk):
    article = get_object_or_404(Article, slug=slug)
    if not (request.user.is_staff or article.author == request.user):
        messages.error(request, 'You do not have permission to edit this article.')
        return redirect(article.get_absolute_url())

    revision = get_object_or_404(ArticleRevision, pk=revision_pk, article=article)

    ArticleRevision.objects.create(
        article=article,
        title=article.title,
        content=article.content,
        excerpt=article.excerpt,
        edited_by=request.user,
    )
    article.title = revision.title
    article.content = revision.content
    article.excerpt = revision.excerpt
    article.save(update_fields=['title', 'content', 'excerpt', 'updated_at'])
    messages.success(request, f'Restored the version from {revision.created_at:%Y-%m-%d %H:%M}.')
    return redirect(article.get_absolute_url())


@login_required
def question_update(request, pk):
    question = get_object_or_404(Question, pk=pk)
    if not (request.user.is_staff or question.author == request.user):
        messages.error(request, 'You do not have permission to edit this question.')
        return redirect(question.get_absolute_url())

    previously_tagged_ids = set(question.tagged_users.values_list('id', flat=True))

    form = QuestionForm(request.POST or None, instance=question)
    if request.method == 'POST' and form.is_valid():
        question = form.save(commit=False)
        _finalize_richtext(question, form, 'body')
        question.save()
        form.save_m2m()
        question.tags.set(_resolve_tags(form.cleaned_data['tags']))
        question.tagged_users.set(form.cleaned_data['tagged_users'])
        newly_tagged = question.tagged_users.exclude(id__in=previously_tagged_ids)
        _notify_tagged_users(request, question, newly_tagged)
        messages.success(request, 'Support request updated.')
        return redirect(question.get_absolute_url())

    context = {
        'content_type': 'qa',
        'form': form,
        'object': question,
    }
    return render(request, 'home/kb_create.html', context)


@login_required
@require_http_methods(['POST'])
def question_delete(request, pk):
    question = get_object_or_404(Question, pk=pk)
    if not (request.user.is_staff or question.author == request.user):
        messages.error(request, 'You do not have permission to delete this support request.')
        return redirect(question.get_absolute_url())
    question.delete()
    messages.success(request, 'Support request deleted.')
    return redirect(f"{reverse('forum-home')}?tab=qa")


@login_required
def report_update(request, pk):
    report = get_object_or_404(Report, pk=pk)
    if not (request.user.is_staff or report.reporter == request.user):
        messages.error(request, 'You do not have permission to edit this report.')
        return redirect(report.get_absolute_url())

    form = ReportForm(request.POST or None, instance=report)
    if request.method == 'POST' and form.is_valid():
        updated_report = form.save(commit=False)
        _finalize_richtext(updated_report, form, 'description')
        if updated_report.status in [Report.STATUS_RESOLVED, Report.STATUS_CLOSED] and not updated_report.resolved_at:
            updated_report.resolved_at = timezone.now()
        if updated_report.status in [Report.STATUS_OPEN, Report.STATUS_IN_PROGRESS]:
            updated_report.resolved_at = None
        updated_report.save()
        form.save_m2m()
        updated_report.tags.set(_resolve_tags(form.cleaned_data['tags']))
        messages.success(request, 'Report updated.')
        return redirect(report.get_absolute_url())

    context = {
        'content_type': 'report',
        'form': form,
        'object': report,
    }
    return render(request, 'home/kb_create.html', context)


@login_required
@require_http_methods(['POST'])
def report_delete(request, pk):
    report = get_object_or_404(Report, pk=pk)
    if not (request.user.is_staff or report.reporter == request.user):
        messages.error(request, 'You do not have permission to delete this report.')
        return redirect(report.get_absolute_url())
    report.delete()
    messages.success(request, 'Report deleted.')
    return redirect(f"{reverse('forum-home')}?tab=reports")


@login_required
def issue_update(request, pk):
    issue = get_object_or_404(IssueLearning, pk=pk)
    if not (request.user.is_staff or issue.reporter == request.user):
        messages.error(request, 'You do not have permission to edit this issue.')
        return redirect(issue.get_absolute_url())

    form = IssueLearningForm(request.POST or None, request.FILES or None, instance=issue)
    if request.method == 'POST' and form.is_valid():
        updated_issue = form.save(commit=False)
        _finalize_richtext(updated_issue, form, 'issue_description')
        _finalize_richtext(updated_issue, form, 'final_corrective_action')
        updated_issue.save()
        form.save_m2m()
        updated_issue.tags.set(_resolve_tags(form.cleaned_data['tags']))
        attachments = form.cleaned_data.get('attachments', [])
        for uploaded_file in attachments:
            IssueLearningAttachment.objects.create(
                issue=updated_issue,
                file=uploaded_file,
                uploaded_by=request.user
            )
        messages.success(request, 'Issue & learning updated.')
        return redirect(issue.get_absolute_url())

    context = {
        'content_type': 'issue',
        'form': form,
        'object': issue,
    }
    return render(request, 'home/kb_create.html', context)


@login_required
@require_http_methods(['POST'])
def issue_delete(request, pk):
    issue = get_object_or_404(IssueLearning, pk=pk)
    if not (request.user.is_staff or issue.reporter == request.user):
        messages.error(request, 'You do not have permission to delete this issue.')
        return redirect(issue.get_absolute_url())
    issue.delete()
    messages.success(request, 'Issue & learning deleted.')
    return redirect(f"{reverse('forum-home')}?tab=issues")


ISSUE_EXCEL_HEADERS = [
    'PROJECT', 'PROBLEM STATEMENT', 'ISSUE DESCRIPTION', 'CLOSURE ACCOUNTABILITY',
    'CLOSURE ACCOUNTABILITY DEPARTMENT', 'LOCATION/AREA', 'ISSUE REPORTED ON',
    'TARGET CLOSURE DATE', 'PRIORITY', 'STATUS', 'CLOSED ON',
    'ROOT CAUSE', 'CORRECTIVE ACTION', 'PREVENTIVE ACTION', 'CAPA OWNER',
    'CAPA TARGET DATE', 'CAPA STATUS', 'CAPA VERIFIED ON', 'CAPA VERIFICATION NOTES',
    'NOTES',
]


def _parse_excel_date(value):
    """openpyxl (data_only=True) hands back a datetime/date for date-formatted
    cells, but a plain string if the source cell was typed as text -- try a
    few common layouts before giving up."""
    import datetime
    if value in (None, ''):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ('%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d', '%d-%b-%Y', '%d-%b-%y'):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


@login_required
def issue_excel_template_download(request):
    import io
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from planner.models import Project as PlannerProject

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Issues & Learnings"

    header_fill = PatternFill(start_color="6B3636", end_color="6B3636", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_side = Side(style='thin', color='4A4A4A')
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col_idx, header in enumerate(ISSUE_EXCEL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = white_bold
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = cell_border
    ws.row_dimensions[1].height = 32

    example_row = [
        'A1477', 'Example: PLC not receiving input signals', 'Describe the issue in detail here.',
        'Jane Doe & John Smith', 'Automation', 'HDPS', '22-08-2026', '29-08-2026',
        'P2', 'Open', '',
        'Why it happened', 'Corrective action once closed.', 'How recurrence is prevented', 'Jane Doe',
        '29-08-2026', 'Open', '', '', 'Optional first note',
    ]
    for col_idx, value in enumerate(example_row, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    widths = [12, 30, 45, 25, 25, 18, 16, 16, 10, 16, 14, 40, 30, 30, 20, 16, 12, 16, 30, 30]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    projects_sheet = wb.create_sheet('Project Codes')
    projects_sheet.cell(row=1, column=1, value='Project Code').font = white_bold
    projects_sheet.cell(row=1, column=1).fill = header_fill
    projects_sheet.cell(row=1, column=2, value='Customer').font = white_bold
    projects_sheet.cell(row=1, column=2).fill = header_fill
    for row_idx, project in enumerate(PlannerProject.objects.order_by('project_id'), start=2):
        projects_sheet.cell(row=row_idx, column=1, value=project.project_id)
        projects_sheet.cell(row=row_idx, column=2, value=project.customer_name)
    projects_sheet.column_dimensions['A'].width = 16
    projects_sheet.column_dimensions['B'].width = 30

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    response = HttpResponse(
        out.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="Issues_Learnings_Template.xlsx"'
    return response


@login_required
def issue_excel_upload(request):
    if request.method == 'POST':
        form = IssueLearningExcelUploadForm(request.POST, request.FILES)
        if form.is_valid():
            import openpyxl
            from planner.models import Project as PlannerProject

            wb = openpyxl.load_workbook(form.cleaned_data['file'], data_only=True)
            ws = wb.worksheets[0]

            header_lookup = {}
            for col_idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
                if cell.value:
                    header_lookup[str(cell.value).strip().upper()] = col_idx

            required = ['PROJECT', 'PROBLEM STATEMENT', 'ISSUE DESCRIPTION', 'ISSUE REPORTED ON']
            missing = [h for h in required if h not in header_lookup]
            if missing:
                messages.error(request, f"Missing required column(s): {', '.join(missing)}.")
                return redirect('kb-issue-excel-upload')

            def cell_value(row, header):
                col_idx = header_lookup.get(header)
                if not col_idx:
                    return None
                value = row[col_idx - 1].value
                return value.strip() if isinstance(value, str) else value

            priority_by_label = {label.upper(): value for value, label in IssueLearning.PRIORITY_CHOICES}
            priority_by_label.update({value.upper(): value for value, _ in IssueLearning.PRIORITY_CHOICES})
            status_by_label = {label.upper(): value for value, label in IssueLearning.STATUS_CHOICES}
            status_by_label.update({value.upper(): value for value, _ in IssueLearning.STATUS_CHOICES})
            capa_status_by_label = {label.upper(): value for value, label in IssueLearning.CAPA_STATUS_CHOICES}
            capa_status_by_label.update({value.upper(): value for value, _ in IssueLearning.CAPA_STATUS_CHOICES})

            created = 0
            row_errors = []
            with transaction.atomic():
                for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
                    if all(c.value in (None, '') for c in row):
                        continue

                    project_code = cell_value(row, 'PROJECT')
                    problem_statement = cell_value(row, 'PROBLEM STATEMENT')
                    issue_description = cell_value(row, 'ISSUE DESCRIPTION')
                    reported_on = _parse_excel_date(cell_value(row, 'ISSUE REPORTED ON'))

                    if not (project_code and problem_statement and issue_description and reported_on):
                        row_errors.append(f"Row {row_idx}: missing a required value (project, problem statement, issue description, or issue reported on date).")
                        continue

                    project = PlannerProject.objects.filter(project_id__iexact=str(project_code).strip()).first()
                    if not project:
                        row_errors.append(f"Row {row_idx}: no project found with code '{project_code}'.")
                        continue

                    priority_raw = cell_value(row, 'PRIORITY')
                    priority = priority_by_label.get(str(priority_raw).strip().upper(), IssueLearning.PRIORITY_P2) if priority_raw else IssueLearning.PRIORITY_P2
                    status_raw = cell_value(row, 'STATUS')
                    status = status_by_label.get(str(status_raw).strip().upper(), IssueLearning.STATUS_OPEN) if status_raw else IssueLearning.STATUS_OPEN
                    capa_status_raw = cell_value(row, 'CAPA STATUS')
                    capa_status = capa_status_by_label.get(str(capa_status_raw).strip().upper(), IssueLearning.CAPA_STATUS_OPEN) if capa_status_raw else IssueLearning.CAPA_STATUS_OPEN

                    issue = IssueLearning.objects.create(
                        project=project,
                        problem_statement=str(problem_statement)[:300],
                        issue_description=str(issue_description),
                        closure_accountability=str(cell_value(row, 'CLOSURE ACCOUNTABILITY') or '')[:300],
                        closure_accountability_department=str(cell_value(row, 'CLOSURE ACCOUNTABILITY DEPARTMENT') or '')[:200],
                        location_area=str(cell_value(row, 'LOCATION/AREA') or '')[:200],
                        issue_reported_on=reported_on,
                        target_closure_date=_parse_excel_date(cell_value(row, 'TARGET CLOSURE DATE')),
                        priority=priority,
                        status=status,
                        closed_on=_parse_excel_date(cell_value(row, 'CLOSED ON')),
                        final_corrective_action=str(cell_value(row, 'CORRECTIVE ACTION') or ''),
                        root_cause=str(cell_value(row, 'ROOT CAUSE') or ''),
                        preventive_action=str(cell_value(row, 'PREVENTIVE ACTION') or ''),
                        capa_owner=str(cell_value(row, 'CAPA OWNER') or '')[:300],
                        capa_target_date=_parse_excel_date(cell_value(row, 'CAPA TARGET DATE')),
                        capa_status=capa_status,
                        capa_verified_on=_parse_excel_date(cell_value(row, 'CAPA VERIFIED ON')),
                        capa_verification_notes=str(cell_value(row, 'CAPA VERIFICATION NOTES') or ''),
                        reporter=request.user,
                    )
                    notes = cell_value(row, 'NOTES')
                    if notes:
                        IssueLearningComment.objects.create(issue=issue, author=request.user, body=str(notes))
                    created += 1

            if created:
                messages.success(request, f"Imported {created} issue(s) & learning(s).")
            for error in row_errors[:10]:
                messages.error(request, error)
            if len(row_errors) > 10:
                messages.error(request, f"...and {len(row_errors) - 10} more row error(s).")
            if created:
                return redirect(f"{reverse('forum-home')}?tab=issues")
        else:
            messages.error(request, 'Please correct the errors below.')
    else:
        form = IssueLearningExcelUploadForm()

    return render(request, 'home/kb_issue_excel_upload.html', {'form': form, 'active_tab': 'issues'})


@login_required
def issue_excel_export(request):
    """
    Export Issues & Learnings (with their CAPA fields) to .xlsx, honoring
    whatever project/status/priority/CAPA-status/search filters are active
    -- reuses _filter_issues() so this always matches what's on screen.
    """
    import io
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from django.utils.html import strip_tags

    issues = _filter_issues(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Issues & Learnings"

    header_fill = PatternFill(start_color="6B3636", end_color="6B3636", fill_type="solid")
    white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrap_align = Alignment(vertical="top", wrap_text=True)
    thin_side = Side(style='thin', color='4A4A4A')
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    headers = ['#'] + ISSUE_EXCEL_HEADERS
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = white_bold
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = cell_border
    ws.row_dimensions[1].height = 32

    date_format = 'DD-MM-YYYY'
    for row_idx, issue in enumerate(issues, start=2):
        notes = '\n'.join(
            f"{comment.created_at:%d-%m-%Y}>> {comment.body}"
            for comment in issue.comments.all()
        )
        values = [
            issue.id,
            issue.project.project_id,
            issue.problem_statement,
            strip_tags(issue.issue_description) if issue.issue_description_is_html else issue.issue_description,
            issue.closure_accountability,
            issue.closure_accountability_department,
            issue.location_area,
            issue.issue_reported_on,
            issue.target_closure_date,
            issue.get_priority_display(),
            issue.get_status_display(),
            issue.closed_on,
            issue.root_cause,
            strip_tags(issue.final_corrective_action) if issue.final_corrective_action_is_html else issue.final_corrective_action,
            issue.preventive_action,
            issue.capa_owner,
            issue.capa_target_date,
            issue.get_capa_status_display(),
            issue.capa_verified_on,
            issue.capa_verification_notes,
            notes,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap_align
            cell.border = cell_border
            if hasattr(value, 'strftime'):
                cell.number_format = date_format

    widths = [6, 12, 30, 45, 25, 25, 18, 16, 16, 10, 16, 14, 30, 30, 30, 20, 16, 12, 16, 30, 30]
    for col_idx, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    response = HttpResponse(
        out.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="Issues_Learnings_Export.xlsx"'
    return response


@login_required
def notifications_list(request):
    notifications = list(
        Notification.objects.filter(recipient=request.user)
        .select_related('actor', 'question', 'report', 'article', 'issue')[:50]
    )
    Notification.objects.filter(recipient=request.user, is_read=False).update(is_read=True)
    return render(request, 'home/notifications.html', {'notifications': notifications})
