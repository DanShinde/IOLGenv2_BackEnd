from datetime import date, timedelta

from django.contrib.auth.models import User
from django.core.exceptions import ValidationError
from django.test import RequestFactory, TestCase
from django.urls import reverse
from django.utils import timezone

from employees.models import Employee
from planner.models import Segment

from .models import (
    RoleMatrix, Skill, SkillBenchmark, SkillMatrix, EmployeeSkill, EmployeeSkillHistory,
    DevelopmentPlan, user_can_manage_employee,
)
from .views import (
    DashboardView, RoleMatrixBenchmarkView, SkillListView, SkillMatrixProfileView,
    designation_benchmark_level_update, safe_json,
)
from .reports import build_team_report_data


def make_skill_matrix(name, segments=None, **kwargs):
    """SkillMatrix.employee is a required OneToOneField to employees.Employee -- this creates
    both in one call, matching the old flat SkillMatrix.objects.create(name=...) call shape
    these tests were originally written against.

    signals.sync_skill_matrix_with_employee already auto-creates a SkillMatrix as soon as the
    Employee is saved (it keeps this app's roster in sync with the shared Employee list), so
    this updates that row with any requested fields (role_matrix, manager, status, user, ...)
    rather than calling SkillMatrix.objects.create() again, which would collide with the
    OneToOneField's unique constraint.

    `segments` sets SkillMatrix.segments (M2M) -- separate from employees.Employee.segment,
    which this app no longer reads for skill-gap purposes.
    """
    employee = Employee.objects.create(name=name, designation='ENGINEER')
    skill_matrix = SkillMatrix.objects.get(employee=employee)
    if kwargs:
        for field, value in kwargs.items():
            setattr(skill_matrix, field, value)
        skill_matrix.save()
    if segments is not None:
        skill_matrix.segments.set(segments)
    return skill_matrix


class GapScoringTests(TestCase):
    def setUp(self):
        self.role = RoleMatrix.objects.create(title='Engineer')
        self.mandatory_skill = Skill.objects.create(name='Python')
        self.mandatory_skill_2 = Skill.objects.create(name='SQL')
        self.optional_skill = Skill.objects.create(name='Public Speaking')

        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.mandatory_skill, required_level=4, is_mandatory=True)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.mandatory_skill_2, required_level=3, is_mandatory=True)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.optional_skill, required_level=3, is_mandatory=False)

        self.employee = make_skill_matrix('Ada Lovelace', role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=self.employee, skill=self.mandatory_skill, actual_level=1)  # gap 3, mandatory
        EmployeeSkill.objects.create(skill_matrix=self.employee, skill=self.mandatory_skill_2, actual_level=3)  # gap 0, mandatory
        EmployeeSkill.objects.create(skill_matrix=self.employee, skill=self.optional_skill, actual_level=0)  # gap 3, optional

    def test_get_skill_gap_data_computes_gap_and_weight(self):
        gaps = {g['skill'].name: g for g in self.employee.get_skill_gap_data()}
        self.assertEqual(gaps['Python']['gap'], 3)
        self.assertEqual(gaps['Python']['weight'], 2)
        self.assertEqual(gaps['Public Speaking']['gap'], 3)
        self.assertEqual(gaps['Public Speaking']['weight'], 1)

    def test_overall_gap_score_weights_mandatory_skills_higher(self):
        # weighted gaps: Python 3*2=6, SQL 0*2=0, Public Speaking 3*1=3; total weight 2+2+1=5
        self.assertAlmostEqual(self.employee.get_overall_gap_score(), 9 / 5)

    def test_skills_met_percentage_is_unweighted(self):
        # 1 of 3 skills met (SQL)
        self.assertEqual(self.employee.get_skills_met_percentage(), round(1 / 3 * 100))

    def test_no_benchmarks_returns_zero(self):
        lone_employee = make_skill_matrix('No Role')
        self.assertEqual(lone_employee.get_overall_gap_score(), 0)
        self.assertEqual(lone_employee.get_skills_met_percentage(), 0)


class SkillHistoryTests(TestCase):
    def setUp(self):
        self.employee = make_skill_matrix('Ada Lovelace')
        self.python = Skill.objects.create(name='Python')
        self.sql = Skill.objects.create(name='SQL')

    def _change_level(self, skill, new_level):
        """Create-then-update so EmployeeSkill.save() creates a history row (skips on create)."""
        es, _ = EmployeeSkill.objects.get_or_create(skill_matrix=self.employee, skill=skill, defaults={'actual_level': 0})
        es.actual_level = new_level
        es.save()
        return es

    def test_skill_with_multiple_changes_returns_ordered_points(self):
        self._change_level(self.python, 3)  # history row: level 3
        es = self._change_level(self.python, 4)  # history row: level 4

        rows = list(EmployeeSkillHistory.objects.filter(skill_matrix=self.employee, skill=self.python).order_by('pk'))
        EmployeeSkillHistory.objects.filter(pk=rows[0].pk).update(evaluated_on='2026-01-10T00:00:00Z')
        EmployeeSkillHistory.objects.filter(pk=rows[1].pk).update(evaluated_on='2026-01-20T00:00:00Z')
        # Current row's last_evaluated coincides with the latest history point, so it collapses
        # into it rather than adding a synthetic third point.
        EmployeeSkill.objects.filter(pk=es.pk).update(last_evaluated=date(2026, 1, 20))

        history = self.employee.get_skill_history()
        self.assertEqual(history['Python'], [
            {'date': '2026-01-10', 'level': 3},
            {'date': '2026-01-20', 'level': 4},
        ])

    def test_skill_with_no_history_rows_still_returns_current_level(self):
        EmployeeSkill.objects.create(skill_matrix=self.employee, skill=self.sql, actual_level=2)
        EmployeeSkill.objects.filter(skill_matrix=self.employee, skill=self.sql).update(last_evaluated=date(2026, 3, 1))

        history = self.employee.get_skill_history()
        self.assertEqual(history['SQL'], [{'date': '2026-03-01', 'level': 2}])

    def test_same_day_history_points_are_deduped_keeping_the_latest(self):
        EmployeeSkillHistory.objects.create(skill_matrix=self.employee, skill=self.python, recorded_level=3)
        EmployeeSkillHistory.objects.create(skill_matrix=self.employee, skill=self.python, recorded_level=4)
        rows = list(EmployeeSkillHistory.objects.filter(skill_matrix=self.employee, skill=self.python).order_by('pk'))
        for row in rows:
            EmployeeSkillHistory.objects.filter(pk=row.pk).update(evaluated_on='2026-01-10T00:00:00Z')
        # No EmployeeSkill row at all for this skill, so only the two history rows are in play.

        history = self.employee.get_skill_history()
        self.assertEqual(history['Python'], [{'date': '2026-01-10', 'level': 4}])


class ViewAccessControlTests(TestCase):
    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin', password='pass12345', is_staff=True)
        self.plain_user = User.objects.create_user('manager', password='pass12345')
        self.role = RoleMatrix.objects.create(title='Engineer')
        self.employee = make_skill_matrix('Ada Lovelace', role_matrix=self.role)

    def test_anonymous_is_redirected_to_login(self):
        response = self.client.get(reverse('skillgap_dashboard'))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse('loginw'), response.url)

        response = self.client.get(reverse('skillgap_employee_list'))
        self.assertEqual(response.status_code, 302)

    def test_logged_in_user_can_view_dashboard(self):
        self.client.force_login(self.plain_user)
        response = self.client.get(reverse('skillgap_dashboard'))
        self.assertEqual(response.status_code, 200)

    def test_non_staff_forbidden_from_delete_view(self):
        self.client.force_login(self.plain_user)
        response = self.client.get(reverse('skillgap_employee_delete', args=[self.employee.id]))
        self.assertEqual(response.status_code, 403)

    def test_staff_can_access_delete_view(self):
        self.client.force_login(self.staff_user)
        response = self.client.get(reverse('skillgap_employee_delete', args=[self.employee.id]))
        self.assertEqual(response.status_code, 200)

    def test_non_staff_forbidden_from_skill_update_endpoint(self):
        self.client.force_login(self.plain_user)
        response = self.client.post(reverse('skillgap_employee_skill_update', args=[self.employee.id]), {})
        self.assertEqual(response.status_code, 403)


class DevelopmentPlanModelTests(TestCase):
    def setUp(self):
        self.employee = make_skill_matrix('Ada Lovelace')
        self.skill = Skill.objects.create(name='Python')

    def _plan(self, **overrides):
        defaults = {'skill_matrix': self.employee, 'skill': self.skill, 'action': 'Take a course'}
        defaults.update(overrides)
        return DevelopmentPlan.objects.create(**defaults)

    def test_is_overdue_true_for_past_target_date_still_open(self):
        plan = self._plan(target_date=timezone.now().date() - timedelta(days=1), status='in_progress')
        self.assertTrue(plan.is_overdue)

    def test_is_overdue_false_for_future_target_date(self):
        plan = self._plan(target_date=timezone.now().date() + timedelta(days=1), status='in_progress')
        self.assertFalse(plan.is_overdue)

    def test_is_overdue_false_when_completed_even_if_past_due(self):
        plan = self._plan(target_date=timezone.now().date() - timedelta(days=1), status='completed')
        self.assertFalse(plan.is_overdue)

    def test_is_overdue_false_without_target_date(self):
        plan = self._plan(target_date=None, status='in_progress')
        self.assertFalse(plan.is_overdue)


class DevelopmentPlanAccessControlTests(TestCase):
    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin', password='pass12345', is_staff=True)
        self.plain_user = User.objects.create_user('manager', password='pass12345')
        self.role = RoleMatrix.objects.create(title='Engineer')
        self.skill = Skill.objects.create(name='Python')
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.skill, required_level=4)
        self.employee = make_skill_matrix('Ada Lovelace', role_matrix=self.role)
        self.plan = DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a course',
        )

    def test_plain_user_can_view_plan_list_and_profile_plans_tab(self):
        self.client.force_login(self.plain_user)
        self.assertEqual(self.client.get(reverse('skillgap_development_plan_list')).status_code, 200)
        self.assertEqual(self.client.get(reverse('skillgap_employee_profile', args=[self.employee.id])).status_code, 200)

    def test_non_staff_forbidden_from_creating_a_plan(self):
        self.client.force_login(self.plain_user)
        response = self.client.get(reverse('skillgap_development_plan_add', args=[self.employee.id]))
        self.assertEqual(response.status_code, 403)

    def test_staff_can_create_update_and_delete_a_plan(self):
        self.client.force_login(self.staff_user)

        response = self.client.post(reverse('skillgap_development_plan_add', args=[self.employee.id]), {
            'skill': self.skill.id,
            'action': 'Pair with a senior dev',
            'status': 'not_started',
        })
        self.assertEqual(response.status_code, 302)
        new_plan = DevelopmentPlan.objects.get(action='Pair with a senior dev')
        self.assertEqual(new_plan.created_by, self.staff_user)

        response = self.client.post(reverse('skillgap_development_plan_update', args=[new_plan.id]), {
            'skill': self.skill.id,
            'action': 'Pair with a senior dev',
            'status': 'in_progress',
        })
        self.assertEqual(response.status_code, 302)
        new_plan.refresh_from_db()
        self.assertEqual(new_plan.status, 'in_progress')

        response = self.client.post(reverse('skillgap_development_plan_delete', args=[new_plan.id]))
        self.assertEqual(response.status_code, 302)
        self.assertFalse(DevelopmentPlan.objects.filter(id=new_plan.id).exists())


class SelfRatingModelTests(TestCase):
    def setUp(self):
        self.employee = make_skill_matrix('Ada Lovelace')
        self.skill = Skill.objects.create(name='Python')

    def test_rating_gap_none_without_self_rating(self):
        es = EmployeeSkill.objects.create(skill_matrix=self.employee, skill=self.skill, actual_level=3)
        self.assertIsNone(es.rating_gap)

    def test_rating_gap_is_self_minus_actual(self):
        es = EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2, self_rated_level=4,
        )
        self.assertEqual(es.rating_gap, 2)

    def test_user_can_manage_employee_staff_always_true(self):
        staff = User.objects.create_user('staff1', password='pass12345', is_staff=True)
        self.assertTrue(user_can_manage_employee(staff, self.employee))

    def test_user_can_manage_employee_true_for_own_manager(self):
        manager_user = User.objects.create_user('mgr1', password='pass12345')
        manager_employee = make_skill_matrix('Manager Mike', user=manager_user)
        self.employee.manager = manager_employee
        self.employee.save()
        self.assertTrue(user_can_manage_employee(manager_user, self.employee))

    def test_user_can_manage_employee_false_for_unrelated_user(self):
        other_user = User.objects.create_user('other1', password='pass12345')
        make_skill_matrix('Someone Else', user=other_user)
        self.assertFalse(user_can_manage_employee(other_user, self.employee))


class SelfRatingWorkflowTests(TestCase):
    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin2', password='pass12345', is_staff=True)
        self.employee_user = User.objects.create_user('employee1', password='pass12345')
        self.manager_user = User.objects.create_user('manager1', password='pass12345')
        self.other_user = User.objects.create_user('other_employee', password='pass12345')

        self.role = RoleMatrix.objects.create(title='Engineer')
        self.skill = Skill.objects.create(name='Python')
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.skill, required_level=4)

        self.manager_employee = make_skill_matrix('Manager Mike', user=self.manager_user)
        self.employee = make_skill_matrix(
            'Ada Lovelace', role_matrix=self.role, user=self.employee_user, manager=self.manager_employee,
        )
        self.unrelated_employee = make_skill_matrix('Someone Else', user=self.other_user)

    def test_employee_can_submit_own_self_rating(self):
        self.client.force_login(self.employee_user)
        response = self.client.post(
            reverse('skillgap_employee_self_rating_update', args=[self.employee.id]),
            {'skill_id': self.skill.id, 'self_rated_level': 3},
        )
        self.assertEqual(response.status_code, 200)
        es = EmployeeSkill.objects.get(skill_matrix=self.employee, skill=self.skill)
        self.assertEqual(es.self_rated_level, 3)
        self.assertEqual(es.rating_status, 'pending')

    def test_employee_cannot_submit_self_rating_for_someone_else(self):
        self.client.force_login(self.employee_user)
        response = self.client.post(
            reverse('skillgap_employee_self_rating_update', args=[self.unrelated_employee.id]),
            {'skill_id': self.skill.id, 'self_rated_level': 3},
        )
        self.assertEqual(response.status_code, 403)

    def test_manager_can_approve_own_reports_self_rating(self):
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        self.client.force_login(self.manager_user)
        response = self.client.post(reverse('skillgap_employee_skill_approve', args=[self.employee.id, self.skill.id]))
        self.assertEqual(response.status_code, 200)
        es = EmployeeSkill.objects.get(skill_matrix=self.employee, skill=self.skill)
        self.assertEqual(es.actual_level, 4)
        self.assertEqual(es.rating_status, 'approved')

    def test_non_manager_cannot_approve_a_rating(self):
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        self.client.force_login(self.other_user)
        response = self.client.post(reverse('skillgap_employee_skill_approve', args=[self.employee.id, self.skill.id]))
        self.assertEqual(response.status_code, 403)

    def test_manager_override_sets_status_overridden(self):
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        self.client.force_login(self.manager_user)
        response = self.client.post(
            reverse('skillgap_employee_skill_update', args=[self.employee.id]),
            {'skill_id': self.skill.id, 'actual_level': 3},
        )
        self.assertEqual(response.status_code, 200)
        es = EmployeeSkill.objects.get(skill_matrix=self.employee, skill=self.skill)
        self.assertEqual(es.actual_level, 3)
        self.assertEqual(es.rating_status, 'overridden')

    def test_my_skills_view_handles_unlinked_account_gracefully(self):
        unlinked_user = User.objects.create_user('unlinked1', password='pass12345')
        self.client.force_login(unlinked_user)
        response = self.client.get(reverse('skillgap_my_skills'))
        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.context['employee'])

    def test_approve_blocked_while_development_plan_is_active(self):
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a course', status='in_progress',
        )
        self.client.force_login(self.manager_user)
        response = self.client.post(reverse('skillgap_employee_skill_approve', args=[self.employee.id, self.skill.id]))
        self.assertEqual(response.status_code, 409)
        self.assertFalse(response.json()['success'])
        es = EmployeeSkill.objects.get(skill_matrix=self.employee, skill=self.skill)
        self.assertEqual(es.actual_level, 2)
        self.assertEqual(es.rating_status, 'pending')

    def test_approve_allowed_once_development_plan_is_completed(self):
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a course', status='completed',
        )
        self.client.force_login(self.manager_user)
        response = self.client.post(reverse('skillgap_employee_skill_approve', args=[self.employee.id, self.skill.id]))
        self.assertEqual(response.status_code, 200)
        es = EmployeeSkill.objects.get(skill_matrix=self.employee, skill=self.skill)
        self.assertEqual(es.actual_level, 4)
        self.assertEqual(es.rating_status, 'approved')

    def test_manager_entering_matching_level_stays_pending_while_plan_active(self):
        """Even if the manager's own entry happens to match the employee's self-rating,
        that must not silently count as approval while a development plan is still open."""
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a course', status='not_started',
        )
        self.client.force_login(self.manager_user)
        response = self.client.post(
            reverse('skillgap_employee_skill_update', args=[self.employee.id]),
            {'skill_id': self.skill.id, 'actual_level': 4},
        )
        self.assertEqual(response.status_code, 200)
        es = EmployeeSkill.objects.get(skill_matrix=self.employee, skill=self.skill)
        self.assertEqual(es.actual_level, 4)
        self.assertEqual(es.rating_status, 'pending')

    def test_cancelled_development_plan_does_not_block_approval(self):
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a course', status='cancelled',
        )
        self.client.force_login(self.manager_user)
        response = self.client.post(reverse('skillgap_employee_skill_approve', args=[self.employee.id, self.skill.id]))
        self.assertEqual(response.status_code, 200)
        es = EmployeeSkill.objects.get(skill_matrix=self.employee, skill=self.skill)
        self.assertEqual(es.rating_status, 'approved')

    def test_has_active_development_plan_false_when_no_plan_exists(self):
        es = EmployeeSkill.objects.create(skill_matrix=self.employee, skill=self.skill, actual_level=2)
        self.assertFalse(es.has_active_development_plan)

    def test_profile_page_shows_blocked_state_for_approve_button(self):
        EmployeeSkill.objects.create(
            skill_matrix=self.employee, skill=self.skill, actual_level=2,
            self_rated_level=4, rating_status='pending',
        )
        DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a course', status='in_progress',
        )
        self.client.force_login(self.manager_user)
        response = self.client.get(reverse('skillgap_employee_profile', args=[self.employee.id]))
        self.assertContains(response, 'Blocked')
        self.assertContains(response, 'disabled')


class RoleMatrixQuerysetTests(TestCase):
    """get_required_benchmarks()/get_missing_benchmarks() must always return a QuerySet,
    never a plain list, since callers chain .values_list()/.filter() onto the result."""

    def test_get_required_benchmarks_is_chainable_without_a_role(self):
        employee = make_skill_matrix('No Role')
        skill_ids = employee.get_required_benchmarks().values_list('skill_id', flat=True)
        self.assertEqual(list(skill_ids), [])

    def test_get_missing_benchmarks_is_chainable_without_a_role(self):
        employee = make_skill_matrix('No Role')
        self.assertEqual(list(employee.get_missing_benchmarks()), [])


class SegmentSkillScopingTests(TestCase):
    """Two-tier skill mapping: General skills apply to every employee in the designation
    regardless of segment; Segment-Specific skills only apply when one of the employee's
    assigned SkillMatrix.segments matches the skill's segment. An employee can be assigned
    more than one segment at once (e.g. someone covering both HDPS and SRM)."""

    def setUp(self):
        self.hdps = Segment.objects.create(name='HDPS')
        self.srm = Segment.objects.create(name='SRM')
        self.vrc = Segment.objects.create(name='VRC')
        self.role = RoleMatrix.objects.create(title='Automation Engineer')

        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.hdps_skill = Skill.objects.create(name='HDPS Controls', scope=Skill.SCOPE_SEGMENT, segment=self.hdps)
        self.srm_skill = Skill.objects.create(name='SRM Robotics', scope=Skill.SCOPE_SEGMENT, segment=self.srm)
        self.vrc_skill = Skill.objects.create(name='VRC Sortation', scope=Skill.SCOPE_SEGMENT, segment=self.vrc)

        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.general_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.hdps_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.srm_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.vrc_skill, required_level=4)

    def test_hdps_employee_sees_general_and_hdps_skills_only(self):
        rahul = make_skill_matrix('Rahul', segments=[self.hdps], role_matrix=self.role)
        skill_names = {b.skill.name for b in rahul.get_required_benchmarks()}
        self.assertEqual(skill_names, {'PLC Programming', 'HDPS Controls'})

    def test_srm_employee_sees_general_and_srm_skills_only(self):
        priya = make_skill_matrix('Priya', segments=[self.srm], role_matrix=self.role)
        skill_names = {b.skill.name for b in priya.get_required_benchmarks()}
        self.assertEqual(skill_names, {'PLC Programming', 'SRM Robotics'})

    def test_employee_with_no_segments_sees_only_general_skills(self):
        manager = make_skill_matrix('No Segment Manager', segments=[], role_matrix=self.role)
        skill_names = {b.skill.name for b in manager.get_required_benchmarks()}
        self.assertEqual(skill_names, {'PLC Programming'})

    def test_employee_working_across_two_segments_sees_both(self):
        """The actual scenario this was added for: someone who works across HDPS and SRM
        should be evaluated on both segments' skills, not just one."""
        arjun = make_skill_matrix('Arjun', segments=[self.hdps, self.srm], role_matrix=self.role)
        skill_names = {b.skill.name for b in arjun.get_required_benchmarks()}
        self.assertEqual(skill_names, {'PLC Programming', 'HDPS Controls', 'SRM Robotics'})
        self.assertNotIn('VRC Sortation', skill_names)

    def test_get_missing_benchmarks_excludes_other_segments_skill(self):
        rahul = make_skill_matrix('Rahul', segments=[self.hdps], role_matrix=self.role)
        missing_names = {s.name for s in rahul.get_missing_benchmarks()}
        self.assertEqual(missing_names, {'PLC Programming', 'HDPS Controls'})
        self.assertNotIn('SRM Robotics', missing_names)

    def test_gap_score_ignores_other_segments_skill_even_if_recorded(self):
        rahul = make_skill_matrix('Rahul', segments=[self.hdps], role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=rahul, skill=self.general_skill, actual_level=3)
        EmployeeSkill.objects.create(skill_matrix=rahul, skill=self.hdps_skill, actual_level=4)
        # Recorded despite not being applicable to any of this employee's segments.
        EmployeeSkill.objects.create(skill_matrix=rahul, skill=self.srm_skill, actual_level=0)

        self.assertEqual(rahul.get_skills_met_percentage(), 100)
        self.assertEqual(rahul.get_overall_gap_score(), 0)

    def test_adding_a_segment_changes_applicable_skill_set_live(self):
        rahul = make_skill_matrix('Rahul', segments=[self.hdps], role_matrix=self.role)
        self.assertNotIn('SRM Robotics', {b.skill.name for b in rahul.get_required_benchmarks()})

        rahul.segments.add(self.srm)

        skill_names = {b.skill.name for b in rahul.get_required_benchmarks()}
        self.assertEqual(skill_names, {'PLC Programming', 'HDPS Controls', 'SRM Robotics'})

    def test_removing_a_segment_changes_applicable_skill_set_live(self):
        rahul = make_skill_matrix('Rahul', segments=[self.hdps, self.srm], role_matrix=self.role)
        rahul.segments.remove(self.hdps)

        skill_names = {b.skill.name for b in rahul.get_required_benchmarks()}
        self.assertEqual(skill_names, {'PLC Programming', 'SRM Robotics'})


class NewEmployeeSegmentSeedTests(TestCase):
    """signals.sync_skill_matrix_with_employee auto-creates a SkillMatrix the moment a new,
    active Employee is saved. That SkillMatrix's `segments` should start seeded from
    Employee.segment (a one-time default), not empty -- otherwise a new hire is silently
    evaluated against General skills only until someone remembers to set this by hand."""

    def setUp(self):
        self.hdps = Segment.objects.create(name='HDPS')

    def test_new_employee_with_a_segment_seeds_skill_matrix_segments(self):
        employee = Employee.objects.create(name='New Hire', designation='ENGINEER', segment=self.hdps)
        skill_matrix = SkillMatrix.objects.get(employee=employee)
        self.assertEqual(list(skill_matrix.segments.all()), [self.hdps])

    def test_new_employee_without_a_segment_leaves_skill_matrix_segments_empty(self):
        employee = Employee.objects.create(name='New Hire', designation='ENGINEER')
        skill_matrix = SkillMatrix.objects.get(employee=employee)
        self.assertEqual(list(skill_matrix.segments.all()), [])

    def test_reactivating_a_previously_inactive_employee_also_seeds_segments(self):
        """The same 'not in the roster yet' branch handles both a brand-new employee and one
        that was inactive since creation and has now been switched on -- both should seed
        the same way, not just the pure-create path."""
        employee = Employee.objects.create(
            name='Rehire', designation='ENGINEER', segment=self.hdps, is_active=False,
        )
        self.assertFalse(SkillMatrix.objects.filter(employee=employee).exists())

        employee.is_active = True
        employee.save()

        skill_matrix = SkillMatrix.objects.get(employee=employee)
        self.assertEqual(list(skill_matrix.segments.all()), [self.hdps])


class SkillScopeValidationTests(TestCase):
    def setUp(self):
        self.srm = Segment.objects.create(name='SRM')

    def test_segment_specific_skill_without_segment_is_invalid(self):
        skill = Skill(name='SRM Robotics', scope=Skill.SCOPE_SEGMENT)
        with self.assertRaises(ValidationError):
            skill.full_clean()

    def test_general_skill_with_segment_is_invalid(self):
        skill = Skill(name='PLC Programming', scope=Skill.SCOPE_GENERAL, segment=self.srm)
        with self.assertRaises(ValidationError):
            skill.full_clean()

    def test_general_skill_without_segment_is_valid(self):
        skill = Skill(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        skill.full_clean()  # should not raise

    def test_segment_specific_skill_with_segment_is_valid(self):
        skill = Skill(name='SRM Robotics', scope=Skill.SCOPE_SEGMENT, segment=self.srm)
        skill.full_clean()  # should not raise

    def test_other_skill_with_segment_is_invalid(self):
        skill = Skill(name='First Aid', scope=Skill.SCOPE_OTHER, segment=self.srm)
        with self.assertRaises(ValidationError):
            skill.full_clean()

    def test_other_skill_without_segment_is_valid(self):
        skill = Skill(name='First Aid', scope=Skill.SCOPE_OTHER)
        skill.full_clean()  # should not raise


class SkillNameUniquenessTests(TestCase):
    """Skill.name is no longer globally unique -- the same name is allowed once per segment
    (e.g. "Commissioning" for ASRS-4D Robot and a separate "Commissioning" for ASRS-HDPS),
    but still can't be duplicated within the same segment, or duplicated across two
    General/Supplementary skills (both have segment=None, which unique_together alone can't
    catch since Postgres never treats two NULLs as equal)."""

    def setUp(self):
        self.hdps = Segment.objects.create(name='ASRS-HDPS')
        self.robot = Segment.objects.create(name='ASRS-4D Robot')

    def test_same_name_allowed_in_two_different_segments(self):
        Skill.objects.create(name='Commissioning', scope=Skill.SCOPE_SEGMENT, segment=self.hdps)
        second = Skill(name='Commissioning', scope=Skill.SCOPE_SEGMENT, segment=self.robot)
        second.full_clean()  # should not raise
        second.save()
        self.assertEqual(Skill.objects.filter(name='Commissioning').count(), 2)

    def test_same_name_rejected_within_the_same_segment(self):
        Skill.objects.create(name='Commissioning', scope=Skill.SCOPE_SEGMENT, segment=self.hdps)
        duplicate = Skill(name='Commissioning', scope=Skill.SCOPE_SEGMENT, segment=self.hdps)
        with self.assertRaises(ValidationError):
            duplicate.full_clean()

    def test_same_name_rejected_for_two_general_skills(self):
        Skill.objects.create(name='Communication', scope=Skill.SCOPE_GENERAL)
        duplicate = Skill(name='Communication', scope=Skill.SCOPE_GENERAL)
        with self.assertRaises(ValidationError):
            duplicate.full_clean()

    def test_same_name_allowed_for_general_and_a_segment_skill(self):
        Skill.objects.create(name='Commissioning', scope=Skill.SCOPE_GENERAL)
        segment_version = Skill(name='Commissioning', scope=Skill.SCOPE_SEGMENT, segment=self.hdps)
        segment_version.full_clean()  # should not raise

    def test_editing_a_skill_does_not_conflict_with_itself(self):
        skill = Skill.objects.create(name='Commissioning', scope=Skill.SCOPE_GENERAL)
        skill.description = 'Updated description'
        skill.full_clean()  # should not raise -- excludes its own pk from the duplicate check


class DevelopmentPlanFormForRoleLessEmployeeTests(TestCase):
    """Regression test: DevelopmentPlanCreateView/UpdateView crashed with AttributeError
    for any employee with no role_matrix, because get_required_benchmarks() used to
    return a plain list and .values_list() was called on it directly."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin3', password='pass12345', is_staff=True)
        self.employee = make_skill_matrix('No Role Employee')
        self.skill = Skill.objects.create(name='Python')

    def test_create_form_loads_for_employee_without_a_role(self):
        self.client.force_login(self.staff_user)
        response = self.client.get(reverse('skillgap_development_plan_add', args=[self.employee.id]))
        self.assertEqual(response.status_code, 200)

    def test_update_form_loads_for_employee_without_a_role(self):
        self.client.force_login(self.staff_user)
        plan = DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a course',
        )
        response = self.client.get(reverse('skillgap_development_plan_update', args=[plan.id]))
        self.assertEqual(response.status_code, 200)


class ConfirmDeletePageTests(TestCase):
    """Regression tests: several confirm-delete templates referenced context variables
    Django's DeleteView never actually provides (e.g. {{ employee.name }} when only
    {{ object }} is set), and benchmark_confirm_delete.html used a nonexistent
    `object.designation` field inside a {% url %} tag, which raised NoReverseMatch."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin4', password='pass12345', is_staff=True)
        self.client.force_login(self.staff_user)
        self.role = RoleMatrix.objects.create(title='Engineer')
        self.skill = Skill.objects.create(name='Python')
        self.employee = make_skill_matrix('Ada Lovelace', role_matrix=self.role)
        self.benchmark = SkillBenchmark.objects.create(role_matrix=self.role, skill=self.skill, required_level=3)

    def test_employee_delete_confirm_shows_employee_name(self):
        response = self.client.get(reverse('skillgap_employee_delete', args=[self.employee.id]))
        self.assertContains(response, 'Ada Lovelace')

    def test_designation_delete_confirm_shows_role_title(self):
        response = self.client.get(reverse('skillgap_designation_delete', args=[self.role.id]))
        self.assertContains(response, 'Engineer')

    def test_benchmark_delete_confirm_renders_without_error(self):
        response = self.client.get(reverse('skillgap_benchmark_delete', args=[self.benchmark.id]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Engineer')

    def test_employee_card_shows_role_title_not_placeholder(self):
        response = self.client.get(reverse('skillgap_employee_card', args=[self.employee.id]))
        self.assertContains(response, 'Engineer')
        self.assertNotContains(response, 'No designation assigned')

    def test_employee_card_shows_development_plans(self):
        DevelopmentPlan.objects.create(
            skill_matrix=self.employee, skill=self.skill, action='Take a Python course',
            status='in_progress',
        )
        response = self.client.get(reverse('skillgap_employee_card', args=[self.employee.id]))
        self.assertContains(response, 'Development Plans')
        self.assertContains(response, 'Take a Python course')
        self.assertContains(response, 'In Progress')

    def test_employee_card_hides_development_plans_section_when_none(self):
        response = self.client.get(reverse('skillgap_employee_card', args=[self.employee.id]))
        self.assertNotContains(response, 'Development Plans')


class SafeJsonTests(TestCase):
    """A skill/employee name containing "</script>" must not be able to close the
    surrounding <script> block early when embedded via {{ ... |safe }}."""

    def test_escapes_script_closing_tag(self):
        rendered = safe_json(['</script><script>alert(1)</script>'])
        self.assertNotIn('</script>', rendered)
        self.assertIn('\\u003c/script\\u003e', rendered)

    def test_still_valid_json_content_for_plain_values(self):
        rendered = safe_json(['Python', 3])
        self.assertIn('"Python"', rendered)
        self.assertIn('3', rendered)


class MalformedInputTests(TestCase):
    """Non-numeric POST values must return a clean error, not an uncaught ValueError."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin5', password='pass12345', is_staff=True)
        self.employee = make_skill_matrix('Ada Lovelace')
        self.skill = Skill.objects.create(name='Python')

    def test_employee_skill_update_rejects_non_numeric_level(self):
        self.client.force_login(self.staff_user)
        response = self.client.post(
            reverse('skillgap_employee_skill_update', args=[self.employee.id]),
            {'skill_id': self.skill.id, 'actual_level': 'not-a-number'},
        )
        self.assertEqual(response.status_code, 400)

    def test_bulk_skill_update_rejects_non_numeric_level(self):
        self.client.force_login(self.staff_user)
        response = self.client.post(reverse('skillgap_bulk_skill_update'), {
            'employee_ids': [self.employee.id],
            'skill_id': self.skill.id,
            'actual_level': 'not-a-number',
        })
        self.assertEqual(response.status_code, 302)
        self.assertFalse(EmployeeSkill.objects.filter(skill_matrix=self.employee, skill=self.skill).exists())


class TeamReportDataTests(TestCase):
    def setUp(self):
        self.role = RoleMatrix.objects.create(title='Engineer')
        self.other_role = RoleMatrix.objects.create(title='Manager')
        self.mandatory_skill = Skill.objects.create(name='Python')
        self.other_role_only_skill = Skill.objects.create(name='Leadership')

        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.mandatory_skill, required_level=4, is_mandatory=True)
        SkillBenchmark.objects.create(role_matrix=self.other_role, skill=self.other_role_only_skill, required_level=3, is_mandatory=False)

        self.engineer = make_skill_matrix('Ada Lovelace', role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=self.engineer, skill=self.mandatory_skill, actual_level=1)  # gap 3

        self.manager = make_skill_matrix('Grace Hopper', role_matrix=self.other_role)
        EmployeeSkill.objects.create(skill_matrix=self.manager, skill=self.other_role_only_skill, actual_level=3)  # gap 0

    def test_matrix_includes_union_of_skills_across_roles(self):
        data = build_team_report_data(SkillMatrix.objects.all())
        skill_names = {s.name for s in data['skills']}
        self.assertEqual(skill_names, {'Python', 'Leadership'})

    def test_skill_not_benchmarked_for_a_role_is_marked_not_applicable(self):
        data = build_team_report_data(SkillMatrix.objects.all())
        cell = data['matrix'][(self.manager.id, self.mandatory_skill.id)]
        self.assertFalse(cell['applicable'])

    def test_applicable_cell_has_correct_actual_and_required(self):
        data = build_team_report_data(SkillMatrix.objects.all())
        cell = data['matrix'][(self.engineer.id, self.mandatory_skill.id)]
        self.assertEqual(cell, {
            'applicable': True, 'required': 4, 'actual': 1, 'gap': 3,
            'status': 'critical', 'is_mandatory': True,
        })

    def test_summary_totals_match_both_employees(self):
        data = build_team_report_data(SkillMatrix.objects.all())
        self.assertEqual(data['summary']['total_employees'], 2)
        self.assertEqual(data['summary']['total_skills'], 2)
        self.assertEqual(data['summary']['skills_met_pct'], 50)  # 1 of 2 comparisons met

    def test_gap_by_employee_only_includes_employees_with_benchmarks(self):
        unbenched = make_skill_matrix('No Role')
        data = build_team_report_data(SkillMatrix.objects.all())
        names = {e['name'] for e in data['gap_by_employee']}
        self.assertIn('Ada Lovelace', names)
        self.assertNotIn(unbenched.name, names)


class TeamReportSegmentScopingTests(TestCase):
    """build_team_report_data() must not show a Segment-Specific skill as 'applicable' for an
    employee whose segments don't include it, even though a benchmark row exists for their
    designation -- regression coverage for the same bug class as SegmentSkillScopingTests,
    but in the Excel/PDF report's own data-building path."""

    def setUp(self):
        self.hdps = Segment.objects.create(name='HDPS')
        self.srm = Segment.objects.create(name='SRM')
        self.role = RoleMatrix.objects.create(title='Automation Engineer')
        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.srm_skill = Skill.objects.create(name='SRM Robotics', scope=Skill.SCOPE_SEGMENT, segment=self.srm)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.general_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.srm_skill, required_level=4)

        self.hdps_employee = make_skill_matrix('Rahul', segments=[self.hdps], role_matrix=self.role)
        # Recorded despite not being applicable -- e.g. via the unscoped Bulk Skill Update picker.
        EmployeeSkill.objects.create(skill_matrix=self.hdps_employee, skill=self.srm_skill, actual_level=2)

    def test_segment_specific_skill_not_applicable_outside_employees_segments(self):
        data = build_team_report_data(SkillMatrix.objects.all())
        cell = data['matrix'][(self.hdps_employee.id, self.srm_skill.id)]
        self.assertEqual(cell, {'applicable': False})

    def test_general_skill_still_applicable(self):
        data = build_team_report_data(SkillMatrix.objects.all())
        cell = data['matrix'][(self.hdps_employee.id, self.general_skill.id)]
        self.assertTrue(cell['applicable'])


class DashboardSegmentInsightsTests(TestCase):
    """DashboardView's Segment Gap Summary -- built via RequestFactory rather than the test
    Client. The Client would run the full middleware stack, including the pre-existing
    SkillGapGroupRequiredMiddleware 403 gap in [[project_skillgap_test_suite_gaps]]; calling
    the view directly exercises get_context_data() without going through that middleware."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin_dash', password='pass12345', is_staff=True)
        self.hdps = Segment.objects.create(name='HDPS')
        self.srm = Segment.objects.create(name='SRM')
        Segment.objects.create(name='VRC')  # no employees -- must not appear in the summary
        self.role = RoleMatrix.objects.create(title='Automation Engineer')

        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.hdps_skill = Skill.objects.create(name='HDPS Controls', scope=Skill.SCOPE_SEGMENT, segment=self.hdps)
        self.srm_skill = Skill.objects.create(name='SRM Robotics', scope=Skill.SCOPE_SEGMENT, segment=self.srm)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.general_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.hdps_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.srm_skill, required_level=4)

        rahul = make_skill_matrix('Rahul', segments=[self.hdps], role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=rahul, skill=self.general_skill, actual_level=3)  # gap 0
        EmployeeSkill.objects.create(skill_matrix=rahul, skill=self.hdps_skill, actual_level=0)  # gap 4

        priya = make_skill_matrix('Priya', segments=[self.srm], role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=priya, skill=self.general_skill, actual_level=3)  # gap 0
        EmployeeSkill.objects.create(skill_matrix=priya, skill=self.srm_skill, actual_level=2)  # gap 2

        # Works across both segments -- their General-skill gap should count toward both
        # segments' averages; their two Segment-Specific skills should count only once each,
        # toward their own segment.
        arjun = make_skill_matrix('Arjun', segments=[self.hdps, self.srm], role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=arjun, skill=self.general_skill, actual_level=0)  # gap 3
        EmployeeSkill.objects.create(skill_matrix=arjun, skill=self.hdps_skill, actual_level=4)  # gap 0
        EmployeeSkill.objects.create(skill_matrix=arjun, skill=self.srm_skill, actual_level=4)  # gap 0

    def _get_dashboard_context(self):
        request = RequestFactory().get(reverse('skillgap_dashboard'))
        request.user = self.staff_user
        response = DashboardView.as_view()(request)
        return response.context_data

    def test_segment_gap_summary_reflects_general_and_segment_specific_gaps(self):
        context = self._get_dashboard_context()
        by_segment = {row['segment']: row for row in context['segment_gap_data']}

        self.assertEqual(set(by_segment), {'HDPS', 'SRM'})  # VRC has no employees, excluded

        # HDPS: Rahul (general gap0*w2=0, hdps gap4*w2=8) + Arjun (general gap3*w2=6, hdps gap0*w2=0)
        # = weighted_gap 14, weight 8 -> avg 1.75
        self.assertEqual(by_segment['HDPS']['avg_gap'], 1.75)
        self.assertEqual(by_segment['HDPS']['employee_count'], 2)

        # SRM: Priya (general gap0*w2=0, srm gap2*w2=4) + Arjun (general gap3*w2=6, srm gap0*w2=0)
        # = weighted_gap 10, weight 8 -> avg 1.25
        self.assertEqual(by_segment['SRM']['avg_gap'], 1.25)
        self.assertEqual(by_segment['SRM']['employee_count'], 2)

    def test_segment_gap_summary_is_sorted_worst_first(self):
        context = self._get_dashboard_context()
        avg_gaps = [row['avg_gap'] for row in context['segment_gap_data']]
        self.assertEqual(avg_gaps, sorted(avg_gaps, reverse=True))


class ProfileSkillGroupingTests(TestCase):
    """SkillMatrixProfileView groups skill_data into General / per-segment (alphabetical) /
    Other (recorded but not part of the current role+segments benchmark) for the Skills
    Assessment tab's sub-tabs -- built via RequestFactory for the same reason as
    DashboardSegmentInsightsTests (sidesteps the pre-existing Client/middleware 403 gap)."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin_profile', password='pass12345', is_staff=True)
        self.hdps = Segment.objects.create(name='HDPS')
        self.srm = Segment.objects.create(name='SRM')
        self.role = RoleMatrix.objects.create(title='Automation Engineer')

        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.hdps_skill = Skill.objects.create(name='HDPS Controls', scope=Skill.SCOPE_SEGMENT, segment=self.hdps)
        self.srm_skill = Skill.objects.create(name='SRM Robotics', scope=Skill.SCOPE_SEGMENT, segment=self.srm)
        self.unrelated_skill = Skill.objects.create(name='Excel', scope=Skill.SCOPE_GENERAL)  # no benchmark for this role

        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.general_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.hdps_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.srm_skill, required_level=4)

        self.arjun = make_skill_matrix('Arjun', segments=[self.hdps, self.srm], role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=self.arjun, skill=self.general_skill, actual_level=3)  # gap 0
        EmployeeSkill.objects.create(skill_matrix=self.arjun, skill=self.hdps_skill, actual_level=0)  # gap 4
        EmployeeSkill.objects.create(skill_matrix=self.arjun, skill=self.srm_skill, actual_level=4)  # gap 0
        EmployeeSkill.objects.create(skill_matrix=self.arjun, skill=self.unrelated_skill, actual_level=2)  # no benchmark

    def _get_profile_context(self, employee):
        request = RequestFactory().get(reverse('skillgap_employee_profile', args=[employee.pk]))
        request.user = self.staff_user
        response = SkillMatrixProfileView.as_view()(request, pk=employee.pk)
        return response.context_data

    def test_general_group_contains_only_general_skills(self):
        context = self._get_profile_context(self.arjun)
        names = {s['skill_name'] for s in context['general_skill_data']}
        self.assertEqual(names, {'PLC Programming'})

    def test_segment_groups_are_alphabetical_and_scoped(self):
        context = self._get_profile_context(self.arjun)
        segment_names = [g['segment_name'] for g in context['segment_skill_groups']]
        self.assertEqual(segment_names, ['HDPS', 'SRM'])

        hdps_group, srm_group = context['segment_skill_groups']
        self.assertEqual({s['skill_name'] for s in hdps_group['skills']}, {'HDPS Controls'})
        self.assertEqual({s['skill_name'] for s in srm_group['skills']}, {'SRM Robotics'})

    def test_unbenchmarked_skill_goes_to_other_group(self):
        context = self._get_profile_context(self.arjun)
        names = {s['skill_name'] for s in context['extra_skill_data']}
        self.assertEqual(names, {'Excel'})

    def test_group_summaries_match_manual_calculation(self):
        context = self._get_profile_context(self.arjun)
        self.assertEqual(context['general_summary'], {'count': 1, 'met_pct': 100, 'avg_gap': 0.0})

        hdps_group, srm_group = context['segment_skill_groups']
        self.assertEqual(hdps_group['summary'], {'count': 1, 'met_pct': 0, 'avg_gap': 4.0})
        self.assertEqual(srm_group['summary'], {'count': 1, 'met_pct': 100, 'avg_gap': 0.0})

        # No required_level -> gap is None -> nothing comparable, met_pct/avg_gap stay None
        # rather than misleadingly reporting 0%/0 gap for a skill with nothing to compare to.
        self.assertEqual(context['extra_summary'], {'count': 1, 'met_pct': None, 'avg_gap': None})

    def test_employee_with_single_segment_only_shows_that_segment_group(self):
        rahul = make_skill_matrix('Rahul', segments=[self.hdps], role_matrix=self.role)
        context = self._get_profile_context(rahul)
        segment_names = [g['segment_name'] for g in context['segment_skill_groups']]
        self.assertEqual(segment_names, ['HDPS'])

    def test_employee_with_no_segments_has_no_segment_groups(self):
        manager = make_skill_matrix('No Segment Manager', segments=[], role_matrix=self.role)
        context = self._get_profile_context(manager)
        self.assertEqual(context['segment_skill_groups'], [])


class OtherSegmentBenchmarksModelTests(TestCase):
    """SkillMatrix.get_other_segment_benchmarks() -- the complement of
    get_required_benchmarks()'s segment-specific half."""

    def setUp(self):
        self.https = Segment.objects.create(name='HTTPS')
        self.pallet = Segment.objects.create(name='Pallet Handling')
        self.case_handling = Segment.objects.create(name='Case Handling')
        self.role = RoleMatrix.objects.create(title='Automation Engineer')

        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.https_skill = Skill.objects.create(name='HTTPS Config', scope=Skill.SCOPE_SEGMENT, segment=self.https)
        self.pallet_skill = Skill.objects.create(name='Pallet Sortation', scope=Skill.SCOPE_SEGMENT, segment=self.pallet)
        self.case_skill = Skill.objects.create(name='Case Handling Logic', scope=Skill.SCOPE_SEGMENT, segment=self.case_handling)

        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.general_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.https_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.pallet_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.case_skill, required_level=3)

    def test_other_benchmarks_excludes_assigned_segments_and_general(self):
        rahul = make_skill_matrix('Rahul', segments=[self.https, self.pallet], role_matrix=self.role)
        other_names = {b.skill.name for b in rahul.get_other_segment_benchmarks()}
        self.assertEqual(other_names, {'Case Handling Logic'})

    def test_other_benchmarks_empty_when_all_role_segments_assigned(self):
        full = make_skill_matrix('Full Coverage', segments=[self.https, self.pallet, self.case_handling], role_matrix=self.role)
        self.assertEqual(list(full.get_other_segment_benchmarks()), [])

    def test_other_benchmarks_is_everything_segment_specific_when_no_segments_assigned(self):
        none_assigned = make_skill_matrix('No Segments', segments=[], role_matrix=self.role)
        other_names = {b.skill.name for b in none_assigned.get_other_segment_benchmarks()}
        self.assertEqual(other_names, {'HTTPS Config', 'Pallet Sortation', 'Case Handling Logic'})

    def test_other_benchmarks_empty_without_a_role(self):
        no_role = make_skill_matrix('No Role')
        self.assertEqual(list(no_role.get_other_segment_benchmarks()), [])


class OtherScopeSkillTests(TestCase):
    """Skill.scope='other' (cross-segment) behaves like General for required-set purposes --
    unconditional, never gated by segment assignment -- but is never returned by
    get_other_segment_benchmarks() (that method is specifically about Segment-Specific
    skills for unassigned segments; an Other skill has no segment to be "unassigned" from)."""

    def setUp(self):
        self.hdps = Segment.objects.create(name='HDPS')
        self.role = RoleMatrix.objects.create(title='Automation Engineer')
        self.other_skill = Skill.objects.create(name='First Aid', scope=Skill.SCOPE_OTHER)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.other_skill, required_level=2)

    def test_other_scope_skill_required_even_with_no_segments_assigned(self):
        employee = make_skill_matrix('No Segments', segments=[], role_matrix=self.role)
        self.assertIn(self.other_skill, [b.skill for b in employee.get_required_benchmarks()])

    def test_other_scope_skill_required_regardless_of_which_segments_are_assigned(self):
        employee = make_skill_matrix('Has HDPS', segments=[self.hdps], role_matrix=self.role)
        self.assertIn(self.other_skill, [b.skill for b in employee.get_required_benchmarks()])

    def test_other_scope_skill_never_appears_in_other_segment_benchmarks(self):
        employee = make_skill_matrix('No Segments', segments=[], role_matrix=self.role)
        self.assertNotIn(self.other_skill, [b.skill for b in employee.get_other_segment_benchmarks()])

    def test_other_scope_skill_gap_counts_toward_overall_score(self):
        employee = make_skill_matrix('No Segments', segments=[], role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=employee, skill=self.other_skill, actual_level=0)
        self.assertEqual(employee.get_overall_gap_score(), 2)  # required 2, actual 0

    def test_profile_page_merges_other_scope_skill_into_general_group(self):
        staff = User.objects.create_user('hr_admin_other_scope', password='pass12345', is_staff=True)
        employee = make_skill_matrix('No Segments', segments=[], role_matrix=self.role)

        request = RequestFactory().get(reverse('skillgap_employee_profile', args=[employee.pk]))
        request.user = staff
        response = SkillMatrixProfileView.as_view()(request, pk=employee.pk)

        names = {s['skill_name'] for s in response.context_data['general_skill_data']}
        self.assertIn('First Aid', names)
        row = next(s for s in response.context_data['general_skill_data'] if s['skill_name'] == 'First Aid')
        self.assertEqual(row['skill_scope'], Skill.SCOPE_OTHER)


class ProfileOtherSkillsTests(TestCase):
    """The profile's "Other Skills" tab: this role's segment-specific skills for segments the
    employee ISN'T assigned to -- shown even at level 0 (never recorded), grouped by segment
    with segments the employee has some existing rating in sorted ahead of untouched ones,
    excluded from the page's gap/benchmark totals, and -- the key promise -- the level already
    recorded there carries over the moment that segment is assigned."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin_other', password='pass12345', is_staff=True)
        self.https = Segment.objects.create(name='HTTPS')
        self.pallet = Segment.objects.create(name='Pallet Handling')
        self.case_handling = Segment.objects.create(name='Case Handling')  # no history for Rahul
        self.srm = Segment.objects.create(name='SRM')  # Rahul has a manager-rated level here
        self.robotic = Segment.objects.create(name='Robotic Palletizers')  # Rahul has a self-rating here
        self.role = RoleMatrix.objects.create(title='Automation Engineer')

        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.https_skill = Skill.objects.create(name='HTTPS Config', scope=Skill.SCOPE_SEGMENT, segment=self.https)
        self.pallet_skill = Skill.objects.create(name='Pallet Sortation', scope=Skill.SCOPE_SEGMENT, segment=self.pallet)
        self.case_skill = Skill.objects.create(name='Case Handling Logic', scope=Skill.SCOPE_SEGMENT, segment=self.case_handling)
        self.srm_skill = Skill.objects.create(name='SRM Robotics', scope=Skill.SCOPE_SEGMENT, segment=self.srm)
        self.robotic_skill = Skill.objects.create(name='Robotic Palletizing', scope=Skill.SCOPE_SEGMENT, segment=self.robotic)

        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.general_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.https_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.pallet_skill, required_level=4)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.case_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.srm_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.robotic_skill, required_level=3)

        self.rahul = make_skill_matrix('Rahul', segments=[self.https, self.pallet], role_matrix=self.role)
        EmployeeSkill.objects.create(skill_matrix=self.rahul, skill=self.general_skill, actual_level=3)
        EmployeeSkill.objects.create(skill_matrix=self.rahul, skill=self.https_skill, actual_level=4)
        EmployeeSkill.objects.create(skill_matrix=self.rahul, skill=self.pallet_skill, actual_level=4)
        # Case Handling Logic: never touched at all -- no EmployeeSkill row.
        EmployeeSkill.objects.create(skill_matrix=self.rahul, skill=self.srm_skill, actual_level=2)  # manager-rated
        EmployeeSkill.objects.create(skill_matrix=self.rahul, skill=self.robotic_skill, actual_level=0, self_rated_level=3)  # self-rated only

    def _get_profile_context(self, employee):
        request = RequestFactory().get(reverse('skillgap_employee_profile', args=[employee.pk]))
        request.user = self.staff_user
        response = SkillMatrixProfileView.as_view()(request, pk=employee.pk)
        return response.context_data

    def test_other_skills_groups_cover_every_unassigned_segment_with_a_benchmark(self):
        context = self._get_profile_context(self.rahul)
        segment_names = {g['segment_name'] for g in context['other_segment_groups']}
        self.assertEqual(segment_names, {'Case Handling', 'SRM', 'Robotic Palletizers'})

    def test_never_recorded_other_skill_still_appears_at_level_zero(self):
        context = self._get_profile_context(self.rahul)
        case_group = next(g for g in context['other_segment_groups'] if g['segment_name'] == 'Case Handling')
        skill = case_group['skills'][0]
        self.assertEqual(skill['skill_name'], 'Case Handling Logic')
        self.assertEqual(skill['actual_level'], 0)
        self.assertEqual(skill['required_level'], 3)
        self.assertEqual(skill['status'], 'not_required')
        self.assertIsNone(skill['gap'])  # not counted toward any gap total

    def test_other_skills_excluded_from_page_level_gap_and_met_totals(self):
        context = self._get_profile_context(self.rahul)
        # Only General + HTTPS + Pallet Handling (all gap 0) count -- Case/SRM/Robotic don't.
        self.assertEqual(context['skills_met'], 100)
        self.assertEqual(context['overall_gap'], 0)

    def test_other_segment_groups_sort_segments_with_history_first_then_alphabetical(self):
        context = self._get_profile_context(self.rahul)
        segment_names = [g['segment_name'] for g in context['other_segment_groups']]
        # Robotic Palletizers (self-rated) and SRM (manager-rated) both have history and sort
        # alphabetically ahead of each other; Case Handling (no history) sorts last despite
        # "Case Handling" being alphabetically first overall.
        self.assertEqual(segment_names, ['Robotic Palletizers', 'SRM', 'Case Handling'])

    def test_recording_a_level_then_assigning_the_segment_carries_it_into_required(self):
        # Record a level for the untouched Case Handling skill directly (staff proactively
        # rating Rahul for a possible future move), same as the "Other Skills" tab would do.
        EmployeeSkill.objects.create(skill_matrix=self.rahul, skill=self.case_skill, actual_level=2)

        context = self._get_profile_context(self.rahul)
        case_group = next(g for g in context['other_segment_groups'] if g['segment_name'] == 'Case Handling')
        self.assertEqual(case_group['skills'][0]['actual_level'], 2)

        self.rahul.segments.add(self.case_handling)

        context = self._get_profile_context(self.rahul)
        self.assertNotIn('Case Handling', {g['segment_name'] for g in context['other_segment_groups']})
        segment_group = next(g for g in context['segment_skill_groups'] if g['segment_name'] == 'Case Handling')
        row = segment_group['skills'][0]
        self.assertEqual(row['skill_name'], 'Case Handling Logic')
        self.assertEqual(row['actual_level'], 2)  # carried over, not reset
        self.assertEqual(row['gap'], 1)  # required 3 - actual 2, now counted

    def test_not_required_skill_data_flattens_groups_preserving_their_sort_order(self):
        """The "Not Currently Required" tab renders one table (not_required_skill_data)
        instead of a sub-header per segment -- this must contain exactly the same rows as
        other_segment_groups + extra_skill_data, in the same relative order, just flattened."""
        context = self._get_profile_context(self.rahul)
        flattened_names = [s['skill_name'] for s in context['not_required_skill_data']]

        expected = []
        for group in context['other_segment_groups']:
            expected.extend(s['skill_name'] for s in group['skills'])
        expected.extend(s['skill_name'] for s in context['extra_skill_data'])

        self.assertEqual(flattened_names, expected)
        # Robotic Palletizers/SRM (history) still come before Case Handling (no history).
        self.assertEqual(flattened_names, ['Robotic Palletizing', 'SRM Robotics', 'Case Handling Logic'])


class RoleMatrixBenchmarkViewTests(TestCase):
    """The Role Matrix page now always lists every catalog skill for every role (General +
    per-segment, alphabetical + Supplementary), whether or not it's currently benchmarked --
    "benchmarking" is just giving a skill a required level, not a separate add step."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin_matrix', password='pass12345', is_staff=True)
        self.https = Segment.objects.create(name='HTTPS')
        self.case_handling = Segment.objects.create(name='Case Handling')
        self.role = RoleMatrix.objects.create(title='Automation Engineer')

        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.unbenchmarked_general_skill = Skill.objects.create(name='Documentation', scope=Skill.SCOPE_GENERAL)
        self.https_skill = Skill.objects.create(name='HTTPS Config', scope=Skill.SCOPE_SEGMENT, segment=self.https)
        self.case_skill = Skill.objects.create(name='Case Handling Logic', scope=Skill.SCOPE_SEGMENT, segment=self.case_handling)
        self.other_skill = Skill.objects.create(name='First Aid', scope=Skill.SCOPE_OTHER)

        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.general_skill, required_level=3)
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.https_skill, required_level=4)
        # Documentation, Case Handling Logic, and First Aid deliberately left unbenchmarked --
        # they must still appear, just with required_level=None.

    def _get_matrix_context(self):
        request = RequestFactory().get(reverse('skillgap_designation_benchmark', args=[self.role.pk]))
        request.user = self.staff_user
        response = RoleMatrixBenchmarkView.as_view()(request, pk=self.role.pk)
        return response.context_data

    def test_every_general_skill_listed_even_if_unbenchmarked(self):
        context = self._get_matrix_context()
        names = {row['skill'].name for row in context['general_rows']}
        self.assertEqual(names, {'PLC Programming', 'Documentation'})

        doc_row = next(r for r in context['general_rows'] if r['skill'].name == 'Documentation')
        self.assertIsNone(doc_row['required_level'])
        plc_row = next(r for r in context['general_rows'] if r['skill'].name == 'PLC Programming')
        self.assertEqual(plc_row['required_level'], 3)

    def test_segment_groups_alphabetical_and_list_every_segment_skill(self):
        context = self._get_matrix_context()
        segment_names = [g['segment_name'] for g in context['segment_groups']]
        self.assertEqual(segment_names, ['Case Handling', 'HTTPS'])  # alphabetical

        case_group = context['segment_groups'][0]
        self.assertEqual(case_group['rows'][0]['skill'].name, 'Case Handling Logic')
        self.assertIsNone(case_group['rows'][0]['required_level'])  # unbenchmarked, still listed

        https_group = context['segment_groups'][1]
        self.assertEqual(https_group['rows'][0]['required_level'], 4)

    def test_other_scope_skills_get_their_own_bucket(self):
        context = self._get_matrix_context()
        self.assertEqual([r['skill'].name for r in context['other_rows']], ['First Aid'])
        self.assertIsNone(context['other_rows'][0]['required_level'])

    def test_benchmark_count_only_counts_actually_benchmarked_skills(self):
        context = self._get_matrix_context()
        self.assertEqual(context['benchmark_count'], 2)  # PLC Programming + HTTPS Config
        self.assertEqual(context['total_skill_count'], 5)  # every catalog skill


class DesignationBenchmarkLevelUpdateTests(TestCase):
    """AJAX endpoint behind the Role Matrix page's inline Required Level inputs: setting a
    level creates/updates the SkillBenchmark; clearing it (empty string) deletes it entirely
    rather than storing a misleading 0."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin_level', password='pass12345', is_staff=True)
        self.plain_user = User.objects.create_user('plain_level', password='pass12345')
        self.role = RoleMatrix.objects.create(title='Automation Engineer')
        self.skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)

    def _post(self, data, user=None):
        request = RequestFactory().post(
            reverse('skillgap_designation_benchmark_level_update', args=[self.role.pk, self.skill.pk]), data=data,
        )
        request.user = user or self.staff_user
        return designation_benchmark_level_update(request, self.role.pk, self.skill.pk)

    def test_setting_a_level_creates_a_benchmark(self):
        response = self._post({'required_level': '4', 'is_mandatory': 'true'})
        self.assertEqual(response.status_code, 200)
        benchmark = SkillBenchmark.objects.get(role_matrix=self.role, skill=self.skill)
        self.assertEqual(benchmark.required_level, 4)
        self.assertTrue(benchmark.is_mandatory)

    def test_setting_a_level_updates_an_existing_benchmark(self):
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.skill, required_level=2, is_mandatory=False)
        self._post({'required_level': '5', 'is_mandatory': 'true'})
        benchmark = SkillBenchmark.objects.get(role_matrix=self.role, skill=self.skill)
        self.assertEqual(benchmark.required_level, 5)
        self.assertTrue(benchmark.is_mandatory)

    def test_clearing_the_level_deletes_the_benchmark(self):
        SkillBenchmark.objects.create(role_matrix=self.role, skill=self.skill, required_level=3)
        response = self._post({'required_level': '', 'is_mandatory': 'true'})
        self.assertEqual(response.status_code, 200)
        self.assertFalse(SkillBenchmark.objects.filter(role_matrix=self.role, skill=self.skill).exists())

    def test_level_zero_is_a_real_value_not_treated_as_clearing(self):
        response = self._post({'required_level': '0', 'is_mandatory': 'true'})
        self.assertEqual(response.status_code, 200)
        benchmark = SkillBenchmark.objects.get(role_matrix=self.role, skill=self.skill)
        self.assertEqual(benchmark.required_level, 0)

    def test_out_of_range_level_rejected(self):
        response = self._post({'required_level': '9', 'is_mandatory': 'true'})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(SkillBenchmark.objects.filter(role_matrix=self.role, skill=self.skill).exists())

    def test_non_numeric_level_rejected(self):
        response = self._post({'required_level': 'abc', 'is_mandatory': 'true'})
        self.assertEqual(response.status_code, 400)

    def test_non_staff_forbidden(self):
        response = self._post({'required_level': '3', 'is_mandatory': 'true'}, user=self.plain_user)
        self.assertEqual(response.status_code, 403)
        self.assertFalse(SkillBenchmark.objects.filter(role_matrix=self.role, skill=self.skill).exists())


class SkillCatalogGroupingTests(TestCase):
    """The Skill Catalog page groups every skill into General / per-segment (alphabetical) /
    Supplementary -- the same category structure the Role Matrix page uses."""

    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin_catalog', password='pass12345', is_staff=True)
        self.https = Segment.objects.create(name='HTTPS')
        self.case_handling = Segment.objects.create(name='Case Handling')

        self.general_skill = Skill.objects.create(name='PLC Programming', scope=Skill.SCOPE_GENERAL)
        self.https_skill = Skill.objects.create(name='HTTPS Config', scope=Skill.SCOPE_SEGMENT, segment=self.https)
        self.case_skill = Skill.objects.create(name='Case Handling Logic', scope=Skill.SCOPE_SEGMENT, segment=self.case_handling)
        self.other_skill = Skill.objects.create(name='First Aid', scope=Skill.SCOPE_OTHER)

    def _get_catalog_context(self):
        request = RequestFactory().get(reverse('skillgap_skill_list'))
        request.user = self.staff_user
        response = SkillListView.as_view()(request)
        return response.context_data

    def test_skills_grouped_into_general_segment_and_other(self):
        context = self._get_catalog_context()
        self.assertEqual([s.name for s in context['general_skills']], ['PLC Programming'])
        self.assertEqual([s.name for s in context['other_skills']], ['First Aid'])

        segment_names = [g['segment_name'] for g in context['segment_skill_groups']]
        self.assertEqual(segment_names, ['Case Handling', 'HTTPS'])  # alphabetical
        self.assertEqual(context['segment_skill_groups'][1]['skills'][0].name, 'HTTPS Config')

    def test_total_skill_count(self):
        context = self._get_catalog_context()
        self.assertEqual(context['total_skill_count'], 4)


class TeamReportExportViewTests(TestCase):
    def setUp(self):
        self.staff_user = User.objects.create_user('hr_admin6', password='pass12345', is_staff=True)
        self.plain_user = User.objects.create_user('plain6', password='pass12345')
        role = RoleMatrix.objects.create(title='Engineer')
        skill = Skill.objects.create(name='Python')
        SkillBenchmark.objects.create(role_matrix=role, skill=skill, required_level=4)
        employee = make_skill_matrix('Ada Lovelace', role_matrix=role)
        EmployeeSkill.objects.create(skill_matrix=employee, skill=skill, actual_level=2)

    def test_anonymous_redirected_from_both_exports(self):
        self.assertEqual(self.client.get(reverse('skillgap_team_report_excel')).status_code, 302)
        self.assertEqual(self.client.get(reverse('skillgap_team_report_pdf')).status_code, 302)

    def test_non_staff_forbidden_from_both_exports(self):
        self.client.force_login(self.plain_user)
        self.assertEqual(self.client.get(reverse('skillgap_team_report_excel')).status_code, 403)
        self.assertEqual(self.client.get(reverse('skillgap_team_report_pdf')).status_code, 403)

    def test_staff_gets_a_valid_excel_file(self):
        self.client.force_login(self.staff_user)
        response = self.client.get(reverse('skillgap_team_report_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        from openpyxl import load_workbook
        from io import BytesIO
        wb = load_workbook(BytesIO(response.content))
        self.assertEqual(wb.sheetnames, ['Dashboard', 'Skill Matrix'])
        self.assertEqual(wb['Skill Matrix']['A1'].value, 'Skill')
        self.assertEqual(wb['Skill Matrix']['C1'].value, 'Ada Lovelace')

    def test_staff_gets_a_valid_pdf_file(self):
        self.client.force_login(self.staff_user)
        response = self.client.get(reverse('skillgap_team_report_pdf'))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))
        self.assertGreater(len(response.content), 500)
