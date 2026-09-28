"""Tests for the tools app.

SimpleTestCase throughout: the app has no models, so there is no database to set up and
the whole suite runs in well under a second. Most of these are regression tests for
specific defects in the Flask app this replaces (TextLists/script.py) -- each one is
labelled with what it guards.
"""
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, SimpleTestCase
from django.urls import reverse

from openpyxl import Workbook, load_workbook

from .forms import SclForm, TextListForm
from .scl import SclParams, render_scl
from .textlist import (
    FORMAT_FERRULES,
    FORMAT_IO,
    FORMAT_IO_AND_FERRULES,
    FORMAT_IO_AND_TAG,
    FORMAT_TAG,
    TextListError,
    analyse,
    build_workbook,
    safe_sheet_name,
)


def make_workbook(sheets, name='io-list.xlsx'):
    """Builds an .xlsx in memory. `sheets` maps a sheet name to (headers, rows)."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, (headers, rows) in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(list(headers))
        for row in rows:
            worksheet.append(list(row))
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        name, buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


STANDARD_HEADERS = ['Tag', 'Channel ', 'I/O Address', 'Ferrules']


def simple_upload(rows, sheet='Panel A', headers=None, name='io-list.xlsx'):
    return make_workbook({sheet: (headers or STANDARD_HEADERS, rows)}, name=name)


def read_back(buffer):
    return load_workbook(BytesIO(buffer.getvalue()))


class TextListParsingTests(SimpleTestCase):

    def test_happy_path_layout_matches_legacy(self):
        """Channel numbers along row 1 from column B, 0-based index down column A."""
        upload = simple_upload([
            ['PB_START', 1, 'I0.0', 'F1'],
            ['PB_STOP', 1, 'I0.1', 'F2'],
            ['LS_01', 2, 'I1.0', 'F3'],
            ['SOL_01', 1, 'Q0.0', 'F4'],
        ])
        buffer, summary = build_workbook(upload, FORMAT_IO_AND_TAG)
        book = read_back(buffer)

        self.assertEqual(book.sheetnames, ['InputsList_Panel A', 'OutputsList_Panel A'])
        inputs = book['InputsList_Panel A']
        self.assertIsNone(inputs['A1'].value)          # index corner stays blank
        self.assertEqual(inputs['B1'].value, 1)        # channel numbers are the headers
        self.assertEqual(inputs['C1'].value, 2)
        self.assertEqual(inputs['A2'].value, 0)        # 0-based index column
        self.assertEqual(inputs['A3'].value, 1)
        self.assertEqual(inputs['B2'].value, 'I0.0 PB_START')
        self.assertEqual(inputs['B3'].value, 'I0.1 PB_STOP')
        self.assertEqual(inputs['C2'].value, 'I1.0 LS_01')
        self.assertEqual(summary.total_input_rows, 3)
        self.assertEqual(summary.total_output_rows, 1)

    def test_every_format(self):
        row = [['PB_START', 1, 'I0.0', 'F1']]
        expected = {
            FORMAT_IO_AND_TAG: 'I0.0 PB_START',
            FORMAT_TAG: 'PB_START',
            FORMAT_IO: 'I0.0',
            FORMAT_FERRULES: 'F1',
            FORMAT_IO_AND_FERRULES: 'I0.0 F1',
        }
        for text_format, value in expected.items():
            with self.subTest(text_format=text_format):
                buffer, _ = build_workbook(simple_upload(row), text_format)
                self.assertEqual(read_back(buffer)['InputsList_Panel A']['B2'].value, value)

    def test_blank_tag_rows_are_dropped_and_counted(self):
        summary = analyse(simple_upload([
            ['PB_START', 1, 'I0.0', 'F1'],
            [None, 1, 'I0.1', 'F2'],
        ]), FORMAT_IO_AND_TAG)
        sheet = summary.sheets[0]
        self.assertEqual(sheet.input_rows, 1)
        self.assertEqual(sheet.blank_tag_rows, 1)

    def test_unreadable_address_is_counted_not_raised(self):
        """Defect: the legacy `word[0]` raised on a blank or numeric I/O Address."""
        summary = analyse(simple_upload([
            ['PB_START', 1, 'I0.0', 'F1'],
            ['BAD_BLANK', 1, None, 'F2'],
            ['BAD_NUMBER', 1, 12, 'F3'],
            ['BAD_LETTER', 1, 'X0.0', 'F4'],
        ]), FORMAT_IO_AND_TAG)
        sheet = summary.sheets[0]
        self.assertEqual(sheet.input_rows, 1)
        self.assertEqual(sheet.unusable_address_rows, 3)

    def test_unreadable_channel_is_counted_not_raised(self):
        """Defect: the legacy `max(sheet['Channel '])` blew up on a NaN or text value."""
        summary = analyse(simple_upload([
            ['PB_START', 1, 'I0.0', 'F1'],
            ['NO_CHANNEL', None, 'I0.1', 'F2'],
            ['TEXT_CHANNEL', 'n/a', 'I0.2', 'F3'],
        ]), FORMAT_IO_AND_TAG)
        sheet = summary.sheets[0]
        self.assertEqual(sheet.input_rows, 1)
        self.assertEqual(sheet.unusable_channel_rows, 2)
        self.assertEqual(sheet.max_channel, 1)

    def test_absurd_channel_is_fatal_and_names_the_row(self):
        with self.assertRaises(TextListError) as caught:
            analyse(simple_upload([['PB_START', 9999, 'I0.0', 'F1']]), FORMAT_IO_AND_TAG)
        self.assertIn('9999', str(caught.exception))
        self.assertIn('row 2', str(caught.exception))

    def test_inputs_only_writes_no_outputs_sheet(self):
        """Defect: `if outputs:` tested a defaultdict any lookup had already keyed, so an
        inputs-only file still got a full-width, entirely empty OutputsList sheet."""
        buffer, summary = build_workbook(
            simple_upload([['PB_START', 1, 'I0.0', 'F1']]), FORMAT_IO_AND_TAG,
        )
        self.assertEqual(read_back(buffer).sheetnames, ['InputsList_Panel A'])
        self.assertEqual(summary.sheets[0].outputs_sheet_name, '')
        self.assertIn('Only inputs found -- no OutputsList sheet will be written.',
                      summary.sheets[0].warnings)

    def test_empty_channels_are_reported(self):
        summary = analyse(simple_upload([
            ['A', 1, 'I0.0', 'F1'],
            ['B', 3, 'I0.1', 'F2'],
        ]), FORMAT_IO_AND_TAG)
        self.assertEqual(summary.sheets[0].empty_channels, [2])


class SheetNameTests(SimpleTestCase):
    """Defect: the legacy app passed the source sheet name straight to Excel, which caps
    names at 31 characters and rejects []:*?/\\ ."""

    def test_long_name_is_truncated(self):
        # 24 chars -- legal as a source sheet name, but 35 once 'InputsList_' is prefixed.
        long_name = 'Level 1 Field IO Panel A'
        buffer, _ = build_workbook(
            simple_upload([['A', 1, 'I0.0', 'F1']], sheet=long_name), FORMAT_IO_AND_TAG,
        )
        written = read_back(buffer).sheetnames[0]
        self.assertLessEqual(len(written), 31)
        self.assertTrue(written.startswith('InputsList_'))

    def test_names_truncating_alike_stay_unique(self):
        first = 'Level 1 Field IO Panel Alpha'
        second = 'Level 1 Field IO Panel Bravo'
        upload = make_workbook({
            first: (STANDARD_HEADERS, [['A', 1, 'I0.0', 'F1']]),
            second: (STANDARD_HEADERS, [['B', 1, 'I0.1', 'F2']]),
        })
        names = read_back(build_workbook(upload, FORMAT_IO_AND_TAG)[0]).sheetnames
        self.assertEqual(len(names), 2)
        self.assertEqual(len(set(names)), 2, f'names collided: {names}')
        self.assertTrue(all(len(n) <= 31 for n in names))

    def test_illegal_characters_are_stripped(self):
        name = safe_sheet_name('InputsList', 'Panel[A]:B*C?D/E\\F', set())
        for char in '[]:*?/\\':
            self.assertNotIn(char, name)


class ColumnMatchingTests(SimpleTestCase):

    def test_trailing_space_is_optional(self):
        """The legacy app required 'Channel ' verbatim; headers now match normalised."""
        for header in ('Channel ', 'Channel', 'channel', ' CHANNEL '):
            with self.subTest(header=header):
                headers = ['Tag', header, 'I/O Address', 'Ferrules']
                summary = analyse(
                    simple_upload([['A', 1, 'I0.0', 'F1']], headers=headers),
                    FORMAT_IO_AND_TAG,
                )
                self.assertEqual(summary.sheets[0].input_rows, 1)

    def test_sheet_without_required_columns_is_reported_not_dropped(self):
        """Defect: sheets missing a column were dropped in silence, so people got an
        empty workbook with no idea why."""
        upload = make_workbook({
            'Cover': (['Project', 'Date'], [['ACME', '2026-01-01']]),
            'Panel A': (STANDARD_HEADERS, [['A', 1, 'I0.0', 'F1']]),
        })
        summary = analyse(upload, FORMAT_IO_AND_TAG)
        self.assertEqual(summary.sheet_count, 1)
        self.assertEqual(len(summary.skipped_sheets), 1)
        self.assertEqual(summary.skipped_sheets[0].name, 'Cover')
        self.assertIn('Tag', summary.skipped_sheets[0].reason)

    def test_no_usable_sheet_is_fatal_and_explains_why(self):
        upload = make_workbook({'Cover': (['Project'], [['ACME']])})
        with self.assertRaises(TextListError) as caught:
            analyse(upload, FORMAT_IO_AND_TAG)
        message = str(caught.exception)
        self.assertIn('Cover', message)
        self.assertIn('Tag', message)

    def test_no_io_rows_at_all_is_fatal(self):
        with self.assertRaises(TextListError) as caught:
            analyse(simple_upload([['A', 1, 'X0.0', 'F1']]), FORMAT_IO_AND_TAG)
        self.assertIn('starts with I or Q', str(caught.exception))


class FerrulesTests(SimpleTestCase):
    """Defect: formats 4 and 5 used a Ferrules column that was never checked for."""

    def test_missing_column_is_fatal_for_ferrule_formats(self):
        headers = ['Tag', 'Channel ', 'I/O Address']
        upload = simple_upload([['A', 1, 'I0.0']], headers=headers)
        with self.assertRaises(TextListError) as caught:
            analyse(upload, FORMAT_FERRULES)
        self.assertIn('Ferrules', str(caught.exception))

    def test_missing_column_is_fine_for_other_formats(self):
        headers = ['Tag', 'Channel ', 'I/O Address']
        summary = analyse(simple_upload([['A', 1, 'I0.0']], headers=headers), FORMAT_TAG)
        self.assertEqual(summary.sheets[0].input_rows, 1)

    def test_blank_values_warn_but_still_download(self):
        upload = simple_upload([
            ['A', 1, 'I0.0', 'F1'],
            ['B', 1, 'I0.1', None],
        ])
        buffer, summary = build_workbook(upload, FORMAT_FERRULES)
        self.assertEqual(summary.sheets[0].missing_ferrules_rows, 1)
        self.assertEqual(read_back(buffer)['InputsList_Panel A']['B3'].value, None)


class AnalyseAgreesWithBuildTests(SimpleTestCase):
    """The whole validate-then-download design rests on these two agreeing."""

    def test_counts_match_cells_written(self):
        upload = make_workbook({
            'Panel A': (STANDARD_HEADERS, [
                ['A', 1, 'I0.0', 'F1'], ['B', 2, 'I1.0', 'F2'], ['C', 1, 'Q0.0', 'F3'],
                [None, 1, 'I0.2', 'F4'], ['E', 1, 'X0.0', 'F5'],
            ]),
            'Panel B': (STANDARD_HEADERS, [['F', 1, 'Q0.1', 'F6']]),
        })
        expected = analyse(upload, FORMAT_IO_AND_TAG)
        buffer, actual = build_workbook(upload, FORMAT_IO_AND_TAG)

        self.assertEqual(expected.total_input_rows, actual.total_input_rows)
        self.assertEqual(expected.total_output_rows, actual.total_output_rows)
        self.assertEqual(expected.fingerprint, actual.fingerprint)

        book = read_back(buffer)
        written = 0
        for name in book.sheetnames:
            worksheet = book[name]
            for row in worksheet.iter_rows(min_row=2, min_col=2):
                written += sum(1 for cell in row if cell.value not in (None, ''))
        self.assertEqual(written, actual.total_input_rows + actual.total_output_rows)


def make_raw_workbook(sheets, name='io-list.xlsx'):
    """Builds an .xlsx from literal rows, so a test can place the header wherever it
    likes (the ARTPL template puts a title banner above it)."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        worksheet = workbook.create_sheet(sheet_name)
        for row in rows:
            worksheet.append(list(row))
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(name, buffer.getvalue())


ARTPL_HEADERS = ['Sr.No', 'Equipment Name', 'Tag', 'Signal Type', 'I/O Address',
                 'Panel Number', 'IO Module Name', 'Module Position', 'Channel', 'Pin']
ARTPL_BANNER = ['ARTPL', '', '', 'IO List', '', '', '', '', '', 'Document No- ARTPL/ATM']


class ArtplFormatTests(SimpleTestCase):
    """The ARTPL IO-list template: a title banner above the header, a Channel column
    that is blank or holds a port name, and IO Module Name as the real grouping unit."""

    def test_header_below_a_banner_row_is_found(self):
        upload = make_raw_workbook({'CP01_IOM': [
            ARTPL_BANNER,
            ARTPL_HEADERS,
            [1, 'CP01', 'Ix_CP01_STR', 'DI', 'I0.0', 'CP01', 'IO101', 'PLC1', None, 1],
            [2, 'CP01', 'Ix_CP01_STP', 'DI', 'I0.1', 'CP01', 'IO101', 'PLC1', None, 2],
        ]})
        summary = analyse(upload, FORMAT_IO_AND_TAG)
        sheet = summary.sheets[0]
        self.assertEqual(sheet.header_row, 2)
        self.assertEqual(sheet.input_rows, 2)
        self.assertIn('Header row found on row 2.', sheet.warnings)

    def test_blank_channel_falls_back_to_io_module_name(self):
        upload = make_raw_workbook({'CP01_IOM': [
            ARTPL_BANNER, ARTPL_HEADERS,
            [1, 'CP01', 'Ix_A', 'DI', 'I0.0', 'CP01', 'IO101', 'PLC1', None, 1],
            [2, 'CP01', 'Ix_B', 'DI', 'I0.1', 'CP01', 'IO102', 'PLC1', None, 1],
        ]})
        buffer, summary = build_workbook(upload, FORMAT_TAG)
        sheet = summary.sheets[0]
        self.assertEqual(sheet.group_by, 'IO Module Name')
        self.assertEqual(sheet.columns, ['IO101', 'IO102'])
        worksheet = read_back(buffer)['InputsList_CP01_IOM']
        self.assertEqual(worksheet['B1'].value, 'IO101')
        self.assertEqual(worksheet['C1'].value, 'IO102')

    def test_text_channel_falls_back_to_io_module_name(self):
        """CP01_DRC has Channel values XM/X4/X2 -- port names, not numbers."""
        upload = make_raw_workbook({'CP01_DRC': [
            ARTPL_HEADERS,
            [1, 'OB_1', 'Ix_A', 'DI', 'I3021.4', 'CP01', 'DRC1', '-', 'X4', 'LP4'],
            [2, 'OB_2', 'Ix_B', 'DI', 'I3021.0', 'CP01', None, '-', 'X2', 'LP2'],
        ]})
        summary = analyse(upload, FORMAT_TAG)
        self.assertEqual(summary.sheets[0].group_by, 'IO Module Name')

    def test_module_name_is_carried_down_the_block(self):
        """It is written once per block and left blank on the rows beneath it."""
        upload = make_raw_workbook({'CP01_DRC': [
            ARTPL_HEADERS,
            [1, 'OB_1', 'Ix_A', 'DI', 'I3021.4', 'CP01', 'DRC1', '-', 'X4', 'LP4'],
            [2, 'OB_2', 'Ix_B', 'DI', 'I3021.0', 'CP01', None, '-', 'X2', 'LP2'],
            [3, 'OB_3', 'Ix_C', 'DI', 'I3085.4', 'CP01', 'DRC2', '-', 'X4', 'LP4'],
            [4, 'OB_4', 'Ix_D', 'DI', 'I3085.0', 'CP01', None, '-', 'X2', 'LP2'],
        ]})
        buffer, summary = build_workbook(upload, FORMAT_TAG)
        self.assertEqual(summary.sheets[0].columns, ['DRC1', 'DRC2'])
        worksheet = read_back(buffer)['InputsList_CP01_DRC']
        self.assertEqual([worksheet['B2'].value, worksheet['B3'].value], ['Ix_A', 'Ix_B'])
        self.assertEqual([worksheet['C2'].value, worksheet['C3'].value], ['Ix_C', 'Ix_D'])

    def test_placeholder_addresses_are_kept_and_counted(self):
        """SPARE rows carry 'Qxxxx'. They stay, so the printed strip lines up with the
        physical terminals, but the summary says how many there were."""
        upload = make_raw_workbook({'CP01_DRC': [
            ARTPL_HEADERS,
            [1, 'SPARE', 'SPARE', 'DO', 'Qxxxx', 'CP01', 'DRC1', '-', 'XM', 'LPM'],
            [2, 'OB_1', 'Qx_REAL', 'DO', 'Q2.0', 'CP01', 'DRC1', '-', 'XM', 'RPM'],
        ]})
        buffer, summary = build_workbook(upload, FORMAT_IO_AND_TAG)
        sheet = summary.sheets[0]
        self.assertEqual(sheet.output_rows, 2)
        self.assertEqual(sheet.placeholder_address_rows, 1)
        self.assertTrue(any('placeholder address' in w for w in sheet.warnings))
        values = [c.value for c in read_back(buffer)['OutputsList_CP01_DRC']['B']][1:]
        self.assertIn('Qxxxx SPARE', values)

    def test_rows_are_ordered_by_pin(self):
        upload = make_raw_workbook({'CP01_IOM': [
            ARTPL_HEADERS,
            [1, 'CP01', 'Ix_THIRD', 'DI', 'I0.2', 'CP01', 'IO101', 'PLC1', None, 3],
            [2, 'CP01', 'Ix_FIRST', 'DI', 'I0.0', 'CP01', 'IO101', 'PLC1', None, 1],
            [3, 'CP01', 'Ix_SECOND', 'DI', 'I0.1', 'CP01', 'IO101', 'PLC1', None, 2],
        ]})
        buffer, _ = build_workbook(upload, FORMAT_TAG)
        column = [c.value for c in read_back(buffer)['InputsList_CP01_IOM']['B']][1:]
        self.assertEqual(column, ['Ix_FIRST', 'Ix_SECOND', 'Ix_THIRD'])

    def test_module_with_no_rows_in_a_direction_gets_no_empty_column(self):
        """IO102 is output-only, so it must not leave a blank column on the input strip."""
        upload = make_raw_workbook({'CP01_IOM': [
            ARTPL_HEADERS,
            [1, 'CP01', 'Ix_A', 'DI', 'I0.0', 'CP01', 'IO101', 'PLC1', None, 1],
            [2, 'CP01', 'Qx_B', 'DO', 'Q2.0', 'CP01', 'IO102', 'PLC1', None, 1],
        ]})
        book = read_back(build_workbook(upload, FORMAT_TAG)[0])
        self.assertEqual([c.value for c in book['InputsList_CP01_IOM'][1]][1:], ['IO101'])
        self.assertEqual([c.value for c in book['OutputsList_CP01_IOM'][1]][1:], ['IO102'])

    def test_cover_sheets_are_skipped_with_a_reason(self):
        upload = make_raw_workbook({
            'Revision Control': [[None, 'Template Revision Control'],
                                 [None, 'Revision', 'Auther', 'Date']],
            'CP01_IOM': [ARTPL_BANNER, ARTPL_HEADERS,
                         [1, 'CP01', 'Ix_A', 'DI', 'I0.0', 'CP01', 'IO101', 'PLC1', None, 1]],
        })
        summary = analyse(upload, FORMAT_IO_AND_TAG)
        self.assertEqual(summary.sheet_count, 1)
        self.assertEqual(summary.skipped_sheets[0].name, 'Revision Control')

    def test_numeric_channel_still_wins_over_module_name(self):
        """A legacy file that also happens to carry IO Module Name keeps its old
        channel-per-column behaviour."""
        upload = make_raw_workbook({'Panel A': [
            ['Tag', 'Channel ', 'I/O Address', 'IO Module Name'],
            ['A', 1, 'I0.0', 'IO101'],
            ['B', 2, 'I0.1', 'IO101'],
        ]})
        summary = analyse(upload, FORMAT_TAG)
        self.assertEqual(summary.sheets[0].group_by, 'Channel')
        self.assertEqual(summary.sheets[0].columns, [1, 2])


class SclTests(SimpleTestCase):

    def test_defaults_reproduce_the_legacy_values(self):
        code = render_scl(SclParams())
        self.assertIn('dbNumber_src := 2, //Output DB Number', code)
        self.assertIn('count := 10);', code)
        self.assertIn('dbNumber_src := 1,   //Input DB Number', code)
        self.assertIn('dbNumber_src := 10,', code)
        self.assertIn('dbNumber_dest := 10,', code)

    def test_db_numbers_are_really_parameters(self):
        """Defect: the legacy body hardcoded DB 1 and 2 while the header claimed to
        describe them, and IBytes was accepted then ignored entirely."""
        code = render_scl(SclParams(input_db=11, output_db=12, hmi_io_db=13, input_bytes=7))
        self.assertNotIn('dbNumber_src := 1,', code)
        self.assertNotIn('dbNumber_src := 2,', code)
        self.assertNotIn('dbNumber_dest := 10,', code)
        self.assertIn('// DB Number 11 is used for Inputs', code)
        self.assertIn('// DB Number 12 is used for Outputs', code)
        self.assertIn('// DB Number 13 is used for HMI IO', code)
        self.assertIn('7 bytes', code)

    def test_form_rejects_out_of_range_values(self):
        base = {'output_bytes': 10, 'input_bytes': 5, 'input_db': 1,
                'output_db': 2, 'hmi_io_db': 10}
        self.assertTrue(SclForm(base).is_valid())
        self.assertFalse(SclForm({**base, 'output_bytes': -1}).is_valid())
        self.assertFalse(SclForm({**base, 'input_db': 0}).is_valid())


class FormTests(SimpleTestCase):

    def test_rejects_non_xlsx(self):
        upload = SimpleUploadedFile('list.csv', b'Tag,Channel', content_type='text/csv')
        form = TextListForm({'text_format': FORMAT_IO_AND_TAG}, {'workbook': upload})
        self.assertFalse(form.is_valid())
        self.assertIn('.xlsx', form.errors['workbook'][0])

    def test_rejects_oversized_file(self):
        oversized = SimpleUploadedFile(
            'big.xlsx', b'x' * (TextListForm.MAX_UPLOAD_BYTES + 1),
        )
        form = TextListForm({'text_format': FORMAT_IO_AND_TAG}, {'workbook': oversized})
        self.assertFalse(form.is_valid())
        self.assertIn('limit is', form.errors['workbook'][0])


class ViewTests(SimpleTestCase):

    def setUp(self):
        self.client = Client()

    def test_pages_are_public(self):
        for name in ('tools_home', 'tools_textlist', 'tools_scl'):
            with self.subTest(name=name):
                self.assertEqual(self.client.get(reverse(name)).status_code, 200)

    def test_validate_over_htmx_returns_only_the_partial(self):
        response = self.client.post(
            reverse('tools_textlist'),
            {'workbook': simple_upload([['A', 1, 'I0.0', 'F1']]),
             'text_format': FORMAT_IO_AND_TAG},
            HTTP_HX_REQUEST='true',
        )
        body = response.content.decode()
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('<html', body)
        self.assertIn('name="fingerprint"', body)
        self.assertIn('InputsList_Panel A', body)

    def test_form_never_lets_htmx_intercept_the_download(self):
        """Regression: with `submit` in the form's hx-trigger, HTMX swallowed the download
        submission too and swapped the raw .xlsx bytes into the page as text. The form
        must submit natively to the download URL, and only the Validate button and the
        format select may go over HTMX."""
        import re
        page = self.client.get(reverse('tools_textlist')).content.decode()
        form = re.search(r'<form id="textlistForm"[^>]*>', page).group(0)
        self.assertIn(f'action="{reverse("tools_textlist_download")}"', form)
        trigger = re.search(r'hx-trigger="([^"]*)"', form).group(1)
        self.assertNotIn('submit', trigger)
        self.assertIn('change from:#id_text_format', trigger)

        partial = self.client.post(
            reverse('tools_textlist'),
            {'workbook': simple_upload([['A', 1, 'I0.0', 'F1']]),
             'text_format': FORMAT_IO_AND_TAG},
            HTTP_HX_REQUEST='true',
        ).content.decode()
        # The opening <button ...> tag whose label is "Download .xlsx".
        download = re.search(
            r'<button([^>]*)>\s*<i[^>]*></i>\s*Download \.xlsx', partial, re.S,
        ).group(1)
        self.assertNotIn('hx-', download)
        self.assertNotIn('formaction', download)
        self.assertIn('type="submit"', download)

    def test_download_streams_an_attachment(self):
        response = self.client.post(
            reverse('tools_textlist_download'),
            {'workbook': simple_upload([['A', 1, 'I0.0', 'F1']]),
             'text_format': FORMAT_IO_AND_TAG},
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response['Content-Disposition'].startswith('attachment;'))
        self.assertIn('spreadsheetml', response['Content-Type'])
        self.assertEqual(read_back(BytesIO(response.content)).sheetnames,
                         ['InputsList_Panel A'])

    def test_download_without_a_file_explains_itself(self):
        response = self.client.post(reverse('tools_textlist_download'),
                                    {'text_format': FORMAT_IO_AND_TAG})
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('Content-Disposition', response)
        # The apostrophe comes back HTML-escaped, so match either side of it.
        self.assertIn('re-sent', response.content.decode())

    def test_changed_file_is_caught_by_the_fingerprint(self):
        response = self.client.post(
            reverse('tools_textlist_download'),
            {'workbook': simple_upload([['A', 1, 'I0.0', 'F1']]),
             'text_format': FORMAT_IO_AND_TAG,
             'fingerprint': '123:deadbeef'},
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('Content-Disposition', response)
        self.assertIn('changed since it was checked', response.content.decode())

    def test_bad_upload_renders_an_error_not_a_download(self):
        upload = make_workbook({'Cover': (['Project'], [['ACME']])})
        response = self.client.post(
            reverse('tools_textlist'),
            {'workbook': upload, 'text_format': FORMAT_IO_AND_TAG},
        )
        self.assertEqual(response.status_code, 200)
        self.assertNotIn('Content-Disposition', response)
        self.assertIn('worksheet', response.content.decode())

    def test_tools_is_excluded_from_the_activity_log(self):
        """Tools is public and stateless -- it stores nothing and changes nothing, so it
        is deliberately not logged. Guarded here so the prefix isn't dropped by accident."""
        from activitylog.middleware import SKIP_PREFIXES
        self.assertIn('/tools/', SKIP_PREFIXES)

    def test_scl_download_streams_text(self):
        response = self.client.post(reverse('tools_scl_download'), {
            'output_bytes': 12, 'input_bytes': 5,
            'input_db': 1, 'output_db': 2, 'hmi_io_db': 10,
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response['Content-Disposition'].startswith('attachment;'))
        self.assertIn('count := 12);', response.content.decode())
